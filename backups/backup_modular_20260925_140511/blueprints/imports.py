# -*- coding: utf-8 -*-
import os, io, re, json, time, glob, zipfile, gzip, mimetypes, tempfile
BASE_DIR = os.path.abspath(os.path.dirname(os.path.dirname(__file__)))
from datetime import datetime, timedelta
from collections import Counter
from flask import (Blueprint, request, jsonify, session, redirect, url_for,
                   render_template, current_app, make_response,
                   send_from_directory, send_file, Response, abort)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from sqlalchemy import func
import pandas as pd
from PIL import Image, ImageOps
try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except ImportError:
    pass
from db import db
from models import (Usuario, Proyecto, AppConfig, TokenStore, NucleusData,
                    NucleusHistory, FiltroMaestro, TablaMaestra,
                    ReglaEstadoManual, AccesoProyecto, KpiConfig,
                    HistorialCambios, Tecnico, Cotizacion)
from services import *


bp = Blueprint('imports', __name__)



@bp.route('/api/import/manual_template')
@login_required
def api_import_manual_template():
    """Generates and downloads an Excel template with primary key + manual columns."""
    from flask import make_response
    pid = session.get('current_proyecto_id')
    if not pid:
        return jsonify({'error': 'No project selected'}), 400

    # Get the primary key column name
    config_key = AppConfig.query.filter_by(proyecto_id=pid, clave='primary_key').first()
    pk_name = config_key.valor if config_key else '_key'

    # Get manual columns
    manual_cfg = AppConfig.query.filter_by(proyecto_id=pid, clave='manual_columns').first()
    manual_cols = [mc['nombre'] for mc in json.loads(manual_cfg.valor)] if manual_cfg else []

    if not manual_cols:
        return jsonify({'error': 'No hay columnas manuales configuradas en este proyecto.'}), 400

    # Fetch all existing primary key values
    rows = NucleusData.query.filter_by(proyecto_id=pid).all()

    proy = db.session.get(Proyecto, pid)
    proy_nombre = proy.nombre if proy else ''

    # Plantilla FLM/PEXT: Detalle (9 campos) + Gestión (16 manuales), con contenidos actuales pre-llenados
    if proy_nombre in ('PEXT', 'FLM'):
        detalle_cols = [
            'Fecha de creación (WO Creation date)',
            'Nombre de Site',
            'Autin TT',
            'Departamento',
            'Fault Level',
            'Estado de la tarea (WO State)',
            'Prioridad del Site',
            'Provincia',
            'Distrito',
        ]
        # Sin _key duplicado; el PK (Número de WO) va como primera columna para el import
        all_cols = detalle_cols + manual_cols
        data = {}
        # PK primero para poder hacer match en el import
        data[pk_name] = [r.key_value for r in rows]
        for col in detalle_cols:
            data[col] = [str((json.loads(r.data_json).get(col) or '')) for r in rows]
        for col in manual_cols:
            data[col] = [str((json.loads(r.data_json).get(col) or '')) for r in rows]
        df = pd.DataFrame(data)
        if df.empty and not rows:
            df = pd.DataFrame(columns=[pk_name] + all_cols)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Plantilla')
            ws = writer.sheets['Plantilla']
            # congelar encabezado y autofiltro
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for col_cells in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col_cells)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 38)
            # resaltar cabeceras de Gestión vs Detalle
            from openpyxl.styles import PatternFill, Font
            fill_det = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
            fill_ges = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            for idx, col_name in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=idx)
                cell.font = Font(bold=True, size=9)
                if col_name in detalle_cols:
                    cell.fill = fill_det
                elif col_name in manual_cols:
                    cell.fill = fill_ges
        output.seek(0)
        response = make_response(output.read())
        response.headers['Content-Disposition'] = f'attachment; filename=plantilla_{proy_nombre}.xlsx'
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return response

    # Build DataFrame: pk column pre-filled, manual columns empty
    df_data = {pk_name: [r.key_value for r in rows]}
    for col in manual_cols:
        df_data[col] = [''] * len(rows)
    df = pd.DataFrame(df_data)

    # Write to in-memory Excel
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Plantilla')
        # Auto-adjust column widths
        ws = writer.sheets['Plantilla']
        for col_cells in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)
    output.seek(0)

    response = make_response(output.read())
    response.headers['Content-Disposition'] = 'attachment; filename=plantilla_columnas_manuales.xlsx'
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


@bp.route('/api/import/preview', methods=['POST'])
@login_required
def api_import_preview():
    file = request.files.get('file')
    if not file: return jsonify({'error': 'No file'}), 400
    try:
        if file.filename.endswith('.csv'):
            df = pd.read_csv(file, encoding='utf-8', nrows=5, dtype=str)
        else:
            df = pd.read_excel(file, nrows=5, dtype=str)
        return jsonify({'columns': [str(c).strip() for c in df.columns]})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@bp.route('/api/import/process', methods=['POST'])
@login_required
def api_import_process():
    try:
        pid = session.get('current_proyecto_id')
        if session.get('rol') == 'demo':
            return jsonify({'error': 'Rol DEMO no tiene permisos para realizar importaciones.'}), 403
        import_type = request.form.get('type') or 'base' # 'base' or 'cruce'
        sum_duplicates = request.form.get('sum_duplicates') == 'true'
        sum_type = request.form.get('sum_type', 'number')
        consolidate_date = request.form.get('consolidate_date') == 'true'
        date_column = request.form.get('date_column', '').strip()
        cols_to_keep_str = request.form.get('columns_to_keep', '[]')
        columns_to_keep = json.loads(cols_to_keep_str)
        file_key = request.form.get('file_key', '').strip()

        # Columnas manuales del proyecto (datos editados por gestores)
        manual_cfg = AppConfig.query.filter_by(proyecto_id=pid, clave='manual_columns').first()
        manual_columns_list = []
        if manual_cfg:
            for mc in json.loads(manual_cfg.valor):
                if isinstance(mc, dict):
                    manual_columns_list.append(str(mc.get('nombre', '')).strip())
        manual_columns_list = [c for c in manual_columns_list if c]
        
        file = request.files.get('file')
        if not file: return jsonify({'error': 'No file'}), 400
        
        if file.filename.endswith('.csv'):
            df = pd.read_csv(file, encoding='utf-8', dtype=str)
        else:
            df = pd.read_excel(file, dtype=str)
            
        df.columns = [str(c).strip() for c in df.columns]
        # Tabulator no soporta '.' en nombres de campo (acceso anidado). Se sanean
        # las columnas aquí para que TODO lo que se importa quede guardado limpio.
        df = df.rename(columns=_sane_data_key)
        df.columns = [str(c).strip() for c in df.columns]
        if not file_key:
            pk_cfg = AppConfig.query.filter_by(proyecto_id=pid, clave='primary_key').first()
            candidate = (pk_cfg.valor if pk_cfg else None) or ''
            if candidate in df.columns:
                file_key = candidate
            elif 'id' in df.columns:
                # Sin pk configurada: preferir "id" antes que la primera columna, que
                # puede ser un valor repetido (p.ej. "Contractor" = "COBRA") y entonces
                # TODAS las filas colapsarían en un único registro.
                file_key = 'id'
            else:
                file_key = str(df.columns[0])
        if file_key not in df.columns:
            return jsonify({'error': f'Key {file_key} not found in headers.'}), 400

        # Validación interna: PEXT debe traer la columna CATEGORY con el valor
        # correcto para evitar importar por error datos de otro proceso/mundo.
        # FLM: sin restricción de CATEGORY (acepta cualquier valor).
        proy_act = Proyecto.query.get(pid)
        proy_act_nombre = proy_act.nombre.strip().upper() if proy_act and proy_act.nombre else ''
        category_esperado = None
        if proy_act_nombre == 'PEXT':
            category_esperado = 'O&M PEXT'
        if category_esperado:
            cat_col = next((c for c in df.columns if c.strip().upper() == 'CATEGORY'), None)
            if cat_col is None:
                return jsonify({'error': f'El archivo no contiene la columna CATEGORY. La importación de {proy_act_nombre} requiere la categoría "{category_esperado}".'}), 400
            valores_cat = {str(v).strip() for v in df[cat_col].dropna().tolist()}
            if valores_cat and not valores_cat.issubset({category_esperado}):
                return jsonify({'error': f'La columna CATEGORY debe contener únicamente "{category_esperado}" para {proy_act_nombre}. Valores detectados: {", ".join(sorted(valores_cat))}.'}), 400

        # Modo columnas manuales: por defecto solo llave + columnas manuales del archivo
        if import_type == 'manual_cols' and not columns_to_keep:
            columns_to_keep = [file_key] + [c for c in manual_columns_list if c in df.columns and c != file_key]

        if columns_to_keep:
            valid_cols = [c for c in columns_to_keep if c in df.columns]
            if file_key not in valid_cols: valid_cols.append(file_key)
            if consolidate_date and date_column and date_column in df.columns and date_column not in valid_cols:
                valid_cols.append(date_column)
            df = df[valid_cols]
        df = df.fillna('')

        if consolidate_date and date_column and date_column in df.columns and len(df) > 0:
            df['_temp_date'] = pd.to_datetime(df[date_column], errors='coerce')
            df = df.sort_values(by='_temp_date', ascending=True, na_position='first')
            df = df.drop_duplicates(subset=[file_key], keep='last')
            df = df.drop(columns=['_temp_date'])
        elif sum_duplicates and len(df) > 0:
            # Smart aggregation: Sum numeric columns, last for the rest
            agg_dict = {}
            for col in df.columns:
                if col == file_key: continue
                
                if sum_type == 'soles':
                    # Clean currency formatting before numeric conversion
                    cleaned_col = df[col].astype(str).str.replace(r'[sS]/\.?\s*', '', regex=True).str.replace(',', '')
                    temp_numeric = pd.to_numeric(cleaned_col, errors='coerce')
                else:
                    temp_numeric = pd.to_numeric(df[col], errors='coerce')
                    
                if not temp_numeric.isna().all():
                    df[col] = temp_numeric.fillna(0)
                    agg_dict[col] = 'sum'
                else:
                    agg_dict[col] = 'last'
            
            if agg_dict:
                df = df.groupby(file_key, as_index=False).agg(agg_dict)
        
        filtros = FiltroMaestro.query.filter_by(proyecto_id=pid).all()
        
        # Agrupar reglas por "Clusters" de columnas (Connected Components)
        # Esto permite que reglas para diferentes columnas se sumen con AND
        # Pero reglas que comparten columnas se sumen con OR
        raw_reglas = []
        for f in filtros:
            cols = [c.strip() for c in f.columna.split(',')]
            vals = [v.strip() for v in f.valor.split(',')]
            raw_reglas.append({'cols': set(cols), 'pairs': list(zip(cols, vals))})
            
        clusters = []
        for r in raw_reglas:
            assigned = False
            for group in clusters:
                # Si la regla comparte alguna columna con el grupo, se une a él
                if any(c in group['columns'] for c in r['cols']):
                    group['columns'].update(r['cols'])
                    group['rules'].append(r['pairs'])
                    assigned = True
                    break
            if not assigned:
                clusters.append({'columns': r['cols'], 'rules': [r['pairs']]})
        
        # Consolidar clusters que puedan haberse cruzado después de unirse por partes
        final_clusters = []
        for c in clusters:
            merged = False
            for f in final_clusters:
                if c['columns'] & f['columns']:
                    f['columns'].update(c['columns'])
                    f['rules'].extend(c['rules'])
                    merged = True
                    break
            if not merged:
                final_clusters.append(c)
            
        tablas = TablaMaestra.query.filter_by(proyecto_id=pid).all()
        reglas = []
        for t in tablas:
            t_cols = [c.strip() for c in t.columna_criterio.split(',')]
            t_vals = [v.strip() for v in t.valor_criterio.split(',')]
            reglas.append({
                'condiciones': list(zip(t_cols, t_vals)),
                'nueva_columna': t.nueva_columna,
                'nuevo_valor': t.nuevo_valor
            })

        reglas_manuales = ReglaEstadoManual.query.filter_by(proyecto_id=pid).all()
        for r in reglas_manuales:
            r_cols = [c.strip() for c in r.columna_criterio.split(',')]
            r_vals = [v.strip() for v in r.valor_criterio.split(',')]
            reglas.append({
                'condiciones': list(zip(r_cols, r_vals)),
                'nueva_columna': r.columna_manual,
                'nuevo_valor': r.nuevo_valor
            })

        config_pk = AppConfig.query.filter_by(proyecto_id=pid, clave='primary_key').first()
        system_pk = config_pk.valor if config_pk else None
        
        if import_type == 'base' and not system_pk:
            system_pk = file_key
            db.session.add(AppConfig(proyecto_id=pid, clave='primary_key', valor=system_pk))
            db.session.commit()
            
        # Delta: consultar solo las filas cuyas claves vienen en el archivo,
        # en chunks de 400 (limite de parametros de SQLite) para no traer toda la tabla.
        imported_keys = set()
        for _v in df[file_key].tolist():
            _s = str(_v).strip()
            if _s:
                imported_keys.add(_s)

        existing_records = {}
        if imported_keys:
            key_list = sorted(imported_keys)
            for i in range(0, len(key_list), 400):
                chunk = key_list[i:i + 400]
                chunk_records = NucleusData.query.filter(
                    NucleusData.proyecto_id == pid,
                    NucleusData.key_value.in_(chunk)
                ).all()
                for r in chunk_records:
                    existing_records[r.key_value] = r
        
        # Consolidation Config
        cons_cfg_row = AppConfig.query.filter_by(proyecto_id=pid, clave='consolidation_config').first()
        cons_cfg = json.loads(cons_cfg_row.valor) if cons_cfg_row else {}
        consolidate_on_fail = cons_cfg.get('consolidate_on_filter_fail', False)
        
        updated, added, ignored, consolidated = 0, 0, 0, 0
        rejected_saldo = []
        dynamic_cols = set()
        counter_guardados = 0

        WO_STATE_COL = 'Estado de la tarea (WO State)'
        STATE_TS_COL = 'FECHA CAMBIO ESTADO'
        DISPATCHED_TS_COL = '_fecha_dispatched'
        CANCEL_REJECT_TS_COL = '_fecha_cancel_reject'
        DISPATCHED_STATES = {'dispatched'}
        TERMINAL_STATES = {'canceled', 'rejected'}
        now_str = datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')

        # Campos protegidos: la importación NO pisa los datos editados manualmente.
        # EXCEPCIÓN: en modo 'manual_cols' (Dataper) el archivo ES la fuente de datos manuales.
        protected_fields = set(manual_columns_list)
        protected_fields.update([
            'SERVICIO', 'CIUDAD', 'TECNICO', 'CONTRATA',
            'MOTIVO DE AVERÍA', 'MOTIVO DE AVERIA', 'SOLUCIÓN', 'SOLUCION',
            'LATITUD', 'LONGITUD', 'SE INSTALÓ MUFAS', 'SE INSTALO MUFAS',
            'LATITUD MUFAS', 'LATITUD Mufas', 'LONGITUD MUFAS', 'LONGITUD Mufas',
            'UBICACIÓN DE MUFAS', 'UBICACION DE MUFAS',
            'SISTEMAS', 'SISTEMA',
            'MATERIALES',
            # Campos del Detalle editables en el modal: el import no pisa lo corregido a mano.
            'Fecha de creación (WO Creation date)', 'FECHA DE CREACIÓN (WO CREATION DATE)',
            'Nombre de Site', 'NOMBRE DE SITE',
            'Autin TT', 'AUTIN TT',
            'Departamento', 'DEPARTAMENTO',
            'Fault Level', 'FAULT LEVEL',
            'Estado de la tarea (WO State)', 'ESTADO DE LA TAREA (WO STATE)',
            'Prioridad del Site', 'PRIORIDAD DEL SITE',
            'Provincia', 'PROVINCIA',
            'Distrito', 'DISTRITO'
        ])
        
        for idx, row in df.iterrows():
            row_dict = row.to_dict()
            key_val = str(row_dict.get(file_key, '')).strip()
            if not key_val: continue
            imported_keys.add(key_val)
            
            # --- 1. MERGE CON DATA EXISTENTE ---
            is_new = False
            if key_val in existing_records:
                record = existing_records[key_val]
                current_data = json.loads(record.data_json)
                old_state = str(current_data.get(WO_STATE_COL, '')).strip()
                if import_type == 'manual_cols':
                    current_data.update(row_dict)
                else:
                    for k, v in row_dict.items():
                        if k in protected_fields and k in current_data:
                            continue
                        current_data[k] = v
                new_state = str(current_data.get(WO_STATE_COL, '')).strip()
                new_state_l = new_state.lower()
                if new_state != old_state:
                    # Registrar TODA transición (incluye quedarse vacío), de modo
                    # que el historial nunca pierda cambios de estado por importación.
                    if new_state_l != old_state.lower():
                        current_data[STATE_TS_COL] = now_str
                        dynamic_cols.add(STATE_TS_COL)
                    db.session.add(HistorialCambios(
                        proyecto_id=pid,
                        usuario_id=None,
                        username='IMPORT',
                        key_value=key_val,
                        campo_modificado=WO_STATE_COL,
                        valor_anterior=old_state,
                        valor_nuevo=new_state,
                        fecha=datetime.utcnow()
                    ))
                if new_state_l in DISPATCHED_STATES and not current_data.get(DISPATCHED_TS_COL):
                    current_data[DISPATCHED_TS_COL] = now_str
                if new_state_l in TERMINAL_STATES and not current_data.get(CANCEL_REJECT_TS_COL):
                    current_data[CANCEL_REJECT_TS_COL] = now_str
            else:
                if import_type == 'cruce':
                    schema_config = AppConfig.query.filter_by(proyecto_id=pid, clave='app_schema').first()
                    schema_val = json.loads(schema_config.valor) if schema_config else []
                    if schema_val:
                        continue
                current_data = row_dict.copy()
                init_state = str(current_data.get(WO_STATE_COL, '')).strip().lower()
                if init_state:
                    current_data[STATE_TS_COL] = now_str
                    dynamic_cols.add(STATE_TS_COL)
                    db.session.add(HistorialCambios(
                        proyecto_id=pid,
                        usuario_id=None,
                        username='IMPORT',
                        key_value=key_val,
                        campo_modificado=WO_STATE_COL,
                        valor_anterior='',
                        valor_nuevo=str(current_data.get(WO_STATE_COL, '')).strip(),
                        fecha=datetime.utcnow()
                    ))
                if init_state in DISPATCHED_STATES:
                    current_data[DISPATCHED_TS_COL] = now_str
                if init_state in TERMINAL_STATES:
                    current_data[CANCEL_REJECT_TS_COL] = now_str
                is_new = True

            # --- 2. APLICAR REGLAS Y TABLAS MAESTRAS ---
            for regla in reglas:
                match = True
                for c, v in regla['condiciones']:
                    if str(current_data.get(c, '')) != v:
                        match = False; break
                if match:
                    current_data[regla['nueva_columna']] = regla['nuevo_valor']
                    dynamic_cols.add(regla['nueva_columna'])

            # --- 3. LOGICA DE FILTRO POR CLUSTERS ---
            keep_record = True
            if final_clusters:
                for cluster in final_clusters:
                    # Usamos current_data para evaluar. Si es un cruce, ahora tiene la info base también.
                    match_cluster = False
                    for rule in cluster['rules']:
                        match_rule = True
                        for c, v in rule:
                            val_archivo = str(current_data.get(c, '')).strip().upper()
                            val_filtro = str(v).strip().upper()
                            if val_archivo != val_filtro:
                                match_rule = False
                                break
                        if match_rule:
                            match_cluster = True
                            break
                    if not match_cluster:
                        keep_record = False
                        break
            
            # --- 4. DECISIÓN DE CONSOLIDACIÓN ---
            if not keep_record:
                if consolidate_on_fail and key_val in existing_records:
                    record = existing_records[key_val]
                    # Guardamos el current_data (que contiene las actualizaciones del archivo importado) en el histórico
                    new_hist = NucleusHistory(proyecto_id=pid, key_value=key_val, data_json=safe_json_dumps(current_data))
                    db.session.add(new_hist)
                    db.session.delete(record)
                    consolidated += 1
                    updated += 1
                    del existing_records[key_val]
                else:
                    ignored += 1
                continue
                
            # --- 4b. COMBUSTIBLE: normalizar FECHA y validar saldo cronológico ---
            # Un GASTO importado tampoco puede dejar en negativo el historial
            # del generador; si lo hace, la fila se omite y se reporta.
            if proy_act_nombre == 'COMBUSTIBLE':
                if 'FECHA' in current_data:
                    current_data['FECHA'] = _combustible_fecha_norm(current_data.get('FECHA', ''))
                _gen_imp = str(current_data.get('QR ASIGNADO', '')).strip()
                _mov_imp = str(current_data.get('MOVIMIENTO', '')).strip().upper()
                _gal_imp = _parse_galones(current_data.get('GALONES'))
                if _gen_imp and _mov_imp == 'GASTO' and _gal_imp > 0:
                    _filas_imp = [f for f in _combustible_filas_gen(pid, _gen_imp)
                                  if str(f.get('key')) != str(key_val)]
                    _filas_imp.append({'key': str(key_val),
                                       'fecha': _combustible_fecha_norm(current_data.get('FECHA', '')),
                                       'mov': 'GASTO', 'gal': _gal_imp})
                    _filas_imp.sort(key=lambda f: (_combustible_fecha_ord(f['fecha']), str(f.get('key') or '')))
                    _ok_imp, _info_imp = _combustible_chequear(_filas_imp)
                    if not _ok_imp:
                        rejected_saldo.append(str(key_val))
                        ignored += 1
                        continue

            # --- 5. GUARDAR REGISTRO ACTIVO ---
            if is_new:
                new_record = NucleusData(proyecto_id=pid, key_value=key_val, data_json=safe_json_dumps(current_data))
                db.session.add(new_record)
                existing_records[key_val] = new_record
                added += 1
            else:
                record = existing_records[key_val]
                record.data_json = safe_json_dumps(current_data)
                updated += 1

            # Commit por lotes: cada 500 registros guarda y libera la transaccion,
            # evitando que una importacion grande exceda el timeout del servidor.
            counter_guardados += 1
            if counter_guardados % 500 == 0:
                db.session.commit()
        
        # Absence-based Consolidation (Optimized to avoid SQLite parameter limits)
        absent_consolidated = 0
        if import_type == 'base' and cons_cfg.get('auto_consolidate_missing'):
            # Consulta solo la columna key_value (no el data_json pesado)
            # para comparar las claves existentes contra las del archivo.
            all_existing_keys = [r[0] for r in db.session.query(NucleusData.key_value)
                                 .filter(NucleusData.proyecto_id == pid).all()]
            for kv in all_existing_keys:
                if kv not in imported_keys:
                    rec = NucleusData.query.filter_by(proyecto_id=pid, key_value=kv).first()
                    if rec is None:
                        continue
                    new_hist = NucleusHistory(proyecto_id=pid, key_value=rec.key_value, data_json=rec.data_json)
                    db.session.add(new_hist)
                    db.session.delete(rec)
                    absent_consolidated += 1
                    consolidated += 1

        db.session.commit()
        
        config_schema = AppConfig.query.filter_by(proyecto_id=pid, clave='app_schema').first()
        schema_cols = set(json.loads(config_schema.valor)) if config_schema else set()
        new_schema = schema_cols.union(set(df.columns)).union(dynamic_cols)
        if new_schema != schema_cols:
            if config_schema:
                config_schema.valor = safe_json_dumps(list(new_schema))
            else:
                db.session.add(AppConfig(proyecto_id=pid, clave='app_schema', valor=safe_json_dumps(list(new_schema))))
            db.session.commit()
        return jsonify({
            'success': True,
            'added': added,
            'updated': updated,
            'ignored': ignored,
            'consolidated': consolidated,
            'rejected_saldo': rejected_saldo,
            'pk': system_pk
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500