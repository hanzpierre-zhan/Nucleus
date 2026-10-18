"""Exportaciones de asistencia. El llamador entrega únicamente filas autorizadas."""

import csv
import io
from flask import Response
from services.asistencia import ESTADOS


def exportar(formato, filas, proyectos, inicio):
    if formato == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Asistencia"
        hoja.append(
            ["Fecha", "Proyecto", "Documento", "Personal", "Estado", "Observación"]
        )
        for r in filas:
            hoja.append(
                [
                    r.fecha,
                    proyectos[r.proyecto_id].nombre,
                    r.documento,
                    r.nombre,
                    ESTADOS[r.estado],
                    r.observacion,
                ]
            )
            hoja.cell(hoja.max_row, 1).number_format = "dd/mm/yyyy"
            for col in range(2, 7):
                celda = hoja.cell(hoja.max_row, col)
                celda.data_type = "s"
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="173E76")
        for col, ancho in enumerate([14, 24, 18, 45, 18, 60], 1):
            hoja.column_dimensions[get_column_letter(col)].width = ancho
        for fila in hoja.iter_rows(min_row=2):
            fila[-1].alignment = Alignment(wrap_text=True, vertical="top")
        hoja.freeze_panes = "D2"
        hoja.auto_filter.ref = hoja.dimensions
        out = io.BytesIO()
        libro.save(out)
        return Response(
            out.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=asistencia-{inicio:%Y-%m}.xlsx"
            },
        )
    if formato == "csv":

        def seguro(v):
            v = str(v)
            return (
                "'" + v
                if v.lstrip().startswith(("=", "+", "-", "@"))
                or v.startswith(("\t", "\r", "\n"))
                else v
            )

        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(
            ["Fecha", "Proyecto", "Documento", "Personal", "Estado", "Observación"]
        )
        for r in filas:
            w.writerow(
                [
                    r.fecha.isoformat(),
                    proyectos[r.proyecto_id].nombre,
                    seguro(r.documento),
                    seguro(r.nombre),
                    ESTADOS[r.estado],
                    seguro(r.observacion),
                ]
            )
        return Response(
            "\ufeff" + out.getvalue(),
            mimetype="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f"attachment; filename=asistencia-{inicio:%Y-%m}.csv"
            },
        )
    return None
