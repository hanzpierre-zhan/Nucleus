"""Catálogo LPU compartido: carga inicial única y edición posterior independiente."""
import io,json,unicodedata
from pathlib import Path
from db import db
from models import LpuRegistro,TokenStore
DEFAULT_CODE='8828126757'
FIELDS={'code':100,'unit':100,'description_es':5000,'description_en':5000,'activity':10000}
HEADERS={'code':'Código único','unit':'Unidad','description_es':'Descripción de la partida (Esp)','description_en':'LPU description (Eng)','activity':'Descripción de la actividad | Description of the activity'}

def seed():
    if TokenStore.query.filter_by(clave='lpu_inicial_v1').first():return
    values=json.loads((Path(__file__).resolve().parent.parent/'data/lpu_inicial.json').read_text(encoding='utf-8'))
    existentes={r.code for r in LpuRegistro.query.all()}
    for value in values:
        if value['code'] not in existentes:db.session.add(LpuRegistro(**value,actualizado_por='Carga inicial LPU'))
    db.session.add(TokenStore(clave='lpu_inicial_v1',valor='214 registros · archivo LPU.xlsx'))
    db.session.commit()

def listado():
    return [serialize(row) for row in LpuRegistro.query.order_by(LpuRegistro.code).all()]

def serialize(row):
    return {'id':row.id,**{field:getattr(row,field) for field in FIELDS}}

def validar(data,parcial=False):
    if not isinstance(data,dict) or set(data)-set(FIELDS) or not data:raise ValueError('Campos LPU inválidos.')
    out={}
    for field,value in data.items():
        if not isinstance(value,str) or len(value)>FIELDS[field]:raise ValueError('Valor o tamaño de campo inválido.')
        out[field]=value.strip()
    if ('code' in out and not out['code']) or (not parcial and not out.get('code')):raise ValueError('El código único es obligatorio.')
    if not parcial:
        for field in FIELDS:out.setdefault(field,'')
    return out

def normalizar(value):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(value or '')).lower() if c.isalnum())

def leer_excel(content):
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    try:
        sheet=book.active
        rows=iter(sheet.iter_rows(values_only=True));header=next(rows,())
        aliases={normalizar(label):field for field,label in HEADERS.items()}
        aliases.update({'each':'unit','codigo':'code','codigounico':'code','unidad':'unit'})
        columns={i:aliases[normalizar(value)] for i,value in enumerate(header) if normalizar(value) in aliases}
        if 'code' not in columns.values():raise ValueError('Falta la columna Código único.')
        result=[];codes=set()
        for number,row in enumerate(rows,2):
            if not any(v is not None and str(v).strip() for v in row):continue
            if number>10001:raise ValueError('El archivo supera los 10 000 registros.')
            item={}
            for i,field in columns.items():
                value=row[i] if i<len(row) else None
                if field=='code' and isinstance(value,(int,float)):
                    if isinstance(value,float) and not value.is_integer():raise ValueError(f'Código inválido en fila {number}.')
                    value=str(int(value))
                item[field]=str(value if value is not None else '')
            if not item.get('activity'):item['activity']=item.get('description_es','')
            try:item=validar(item)
            except ValueError as error:raise ValueError(f'Fila {number}: {error}')
            if item['code'] in codes:raise ValueError(f'Código repetido en el archivo: {item["code"]}.')
            codes.add(item['code']);result.append(item)
        if not result:raise ValueError('El archivo no contiene registros.')
        return result
    finally:book.close()

def excel(template=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    book=Workbook();sheet=book.active;sheet.title='LPU';keys=list(HEADERS)
    sheet.append([HEADERS[key] for key in keys]);sheet.freeze_panes='A2'
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='145B67');cell.alignment=Alignment(wrap_text=True)
    sheet.row_dimensions[1].height=32
    for key,width in zip('ABCDE',(22,14,70,70,70)):sheet.column_dimensions[key].width=width
    if not template:
        for value in listado():sheet.append([value[key] for key in keys])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:cell.number_format='@';cell.alignment=Alignment(vertical='top',wrap_text=True);cell.data_type='s'
    sheet.auto_filter.ref=sheet.dimensions
    output=io.BytesIO();book.save(output);output.seek(0);return output
