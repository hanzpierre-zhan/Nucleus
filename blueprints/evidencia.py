# -*- coding: utf-8 -*-
import os, io, re, json, time, glob, zipfile, gzip, mimetypes, tempfile
BASE_DIR = os.path.abspath(os.path.dirname(os.path.dirname(__file__)))
from datetime import datetime, timedelta
from collections import Counter
from flask import (Blueprint, request, jsonify, session, redirect, url_for,
                   render_template, current_app, make_response,
                   send_from_directory, send_file, Response, abort)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from sqlalchemy import func
import pandas as pd
from PIL import Image, ImageOps
try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except ImportError:
    pass
from db import db
from models import (Usuario, Proyecto, AppConfig, TokenStore, NucleusData,
                    NucleusHistory, FiltroMaestro, TablaMaestra,
                    ReglaEstadoManual, AccesoProyecto, KpiConfig,
                    HistorialCambios, Tecnico, Cotizacion)
from services import *


bp = Blueprint('evidencia', __name__)



@bp.route('/api/evidencia/subir', methods=['POST'])
@login_required
def api_evidencia_subir():
    if session.get('rol') == 'demo':
        return jsonify({'error': 'Rol DEMO no tiene permisos para subir evidencia.'}), 403
    pid = session.get('current_proyecto_id')
    key = (request.form.get('key') or '').strip()
    tipo = (request.form.get('tipo') or '').strip().lower()
    if _evidencia_aprobacion_bloquea(pid, key):
        return jsonify({'error': 'Este WO ya fue enviado a aprobación. Solo el personal administrativo puede cambiar su evidencia.'}), 403
    try:
        indice = int(request.form.get('indice'))
    except (TypeError, ValueError):
        return jsonify({'error': 'Índice inválido'}), 400
    file = request.files.get('foto')
    if not key or tipo not in EVIDENCIA_TIPOS + ('comb', 'pex', EVIDENCIA_TIPO_RESGUARDO):
        return jsonify({'error': 'Datos incompletos'}), 400
    if tipo == 'pex':
        max_i = _pext_max()
    elif tipo == EVIDENCIA_TIPO_RESGUARDO:
        max_i = EVIDENCIA_MAX_RESGUARDO
    else:
        max_i = EVIDENCIA_MAX_POR_TIPO
    if indice < 0 or indice >= max_i:
        return jsonify({'error': 'Índice fuera de rango'}), 400
    if file is None or not file.filename:
        return jsonify({'error': 'No se recibió ningún archivo'}), 400
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in EVIDENCIA_EXT_ALLOWED:
        return jsonify({'error': f'Formato no permitido: {ext}. Usa {", ".join(sorted(EVIDENCIA_EXT_ALLOWED))}'}), 400

    # Guardar a un archivo temporal, comprimir y luego mover/subir.
    fd, ruta_tmp = tempfile.mkstemp(suffix=ext)
    os.close(fd)
    try:
        file.save(ruta_tmp)
        nombre = f'{tipo}_{indice}{ext}'
        if evidencia_comprimir(ruta_tmp):
            if not nombre.lower().endswith('.jpg'):
                nueva = ruta_tmp + '.jpg'
                os.rename(ruta_tmp, nueva)
                ruta_tmp = nueva
                nombre = f'{tipo}_{indice}.jpg'

        if evidencia_usa_b2():
            evidencia_eliminar_b2(key, tipo, indice)
            b2_cliente().upload_file(ruta_tmp, current_app.config['B2_BUCKET'], f'{key}/{nombre}')
        else:
            folder = evidencia_folder(pid, key)
            os.makedirs(folder, exist_ok=True)
            evidencia_limpiar_slot(pid, key, tipo, indice)
            os.replace(ruta_tmp, os.path.join(folder, nombre))
            # FLM <-> FLM - ENTEL: replica la foto física en la carpeta del proyecto
            # hermano para que la evidencia sea visible/eliminable desde ambos.
            her = _flm_hermano_id(pid)
            if her is not None:
                her_folder = evidencia_folder(her, key)
                os.makedirs(her_folder, exist_ok=True)
                evidencia_limpiar_slot(her, key, tipo, indice)
                try:
                    import shutil
                    shutil.copy2(os.path.join(folder, nombre), os.path.join(her_folder, nombre))
                except Exception:
                    pass

        # 2do backup: OneDrive personal. Si falla, no interrumpe la subida principal.
        try:
            onedrive_subir(f'{key}/{nombre}', ruta_tmp)
        except Exception as e:
            current_app.logger.warning('OD backup fallo: %s', e)

        url = f'/api/evidencia/foto/{pid}/{secure_filename(str(key))}/{nombre}?v={int(time.time())}'
        return jsonify({'success': True, 'url': url})
    finally:
        if os.path.exists(ruta_tmp):
            try:
                os.remove(ruta_tmp)
            except Exception:
                pass


@bp.route('/api/evidencia/eliminar', methods=['POST'])
@login_required
def api_evidencia_eliminar():
    if session.get('rol') == 'demo':
        return jsonify({'error': 'Rol DEMO no tiene permisos.'}), 403
    pid = session.get('current_proyecto_id')
    data = request.json or {}
    key = (data.get('key') or '').strip()
    tipo = (data.get('tipo') or '').strip().lower()
    if _evidencia_aprobacion_bloquea(pid, key):
        return jsonify({'error': 'Este WO ya fue enviado a aprobación. Solo el personal administrativo puede cambiar su evidencia.'}), 403
    try:
        indice = int(data.get('indice'))
    except (TypeError, ValueError):
        return jsonify({'error': 'Índice inválido'}), 400
    if not key or tipo not in EVIDENCIA_TIPOS + ('comb', 'pex', EVIDENCIA_TIPO_RESGUARDO):
        return jsonify({'error': 'Datos incompletos'}), 400
    if tipo == 'pex':
        max_i = _pext_max()
    elif tipo == EVIDENCIA_TIPO_RESGUARDO:
        max_i = EVIDENCIA_MAX_RESGUARDO
    else:
        max_i = EVIDENCIA_MAX_POR_TIPO
    if indice < 0 or indice >= max_i:
        return jsonify({'error': 'Índice inválido'}), 400
    # Solo admin puede borrar foto de Combustible
    if tipo == 'comb' and session.get('rol') not in ('zeno', 'suport'):
        return jsonify({'error': 'Solo el administrador puede eliminar la foto.'}), 403
    if evidencia_usa_b2():
        evidencia_eliminar_b2(key, tipo, indice)
    else:
        evidencia_limpiar_slot(pid, key, tipo, indice)
        # FLM <-> FLM - ENTEL: limpiar también en la carpeta del proyecto hermano.
        her = _flm_hermano_id(pid)
        if her is not None:
            evidencia_limpiar_slot(her, key, tipo, indice)
    try:
        onedrive_eliminar(key, tipo, indice)
    except Exception as e:
        current_app.logger.warning('OD delete fallo: %s', e)
    return jsonify({'success': True})


@bp.route('/api/evidencia/foto/<int:pid>/<path:key>/<path:nombre>')
@login_required
def api_evidencia_foto(pid, key, nombre):
    cur = session.get('current_proyecto_id')
    # FLM <-> FLM - ENTEL: las fotos subidas desde uno son accesibles desde el otro.
    permitido = (pid == cur) or (_flm_hermano_id(cur) == pid)
    if not permitido:
        return jsonify({'error': 'Acceso denegado'}), 403
    nombre = os.path.basename(nombre)
    if evidencia_usa_b2():
        try:
            obj = b2_cliente().get_object(Bucket=current_app.config['B2_BUCKET'], Key=f'{key}/{nombre}')
            data = obj['Body'].read()
            mt = mimetypes.guess_type(nombre)[0] or 'application/octet-stream'
            resp = Response(data, mimetype=mt)
            resp.headers['Cache-Control'] = 'private, max-age=604800, immutable'
            return resp
        except Exception:
            return jsonify({'error': 'No encontrado'}), 404
    folder = evidencia_folder(pid, key)
    ruta = os.path.join(folder, nombre)
    if not os.path.exists(ruta):
        if '/' not in nombre and '\\' not in nombre and not nombre.startswith('.') and os.path.isdir(folder):
            pre, _ = os.path.splitext(nombre)
            candidatos = [f for f in os.listdir(folder) if f == nombre or f.startswith(pre + '.')]
            if candidatos:
                nombre = sorted(candidatos)[0]
    return send_from_directory(folder, nombre, max_age=604800)


@bp.route('/api/evidencia/zip/<int:pid>/<path:key>')
@login_required
def api_evidencia_zip(pid, key):
    if pid != session.get('current_proyecto_id'):
        return jsonify({'error': 'Acceso denegado'}), 403
    key = (key or '').strip()
    if not key:
        return jsonify({'error': 'Ticket sin clave'}), 400
    # Sanitizar para evitar path traversal
    if '/' in key or '\\' in key or '..' in key:
        return jsonify({'error': 'Clave inválida'}), 400
    zip_buf = io.BytesIO()
    count = 0
    with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        if evidencia_usa_b2():
            try:
                client = b2_cliente()
                bucket = current_app.config['B2_BUCKET']
                resp = client.list_objects_v2(Bucket=bucket, Prefix=f'{key}/')
                for obj in resp.get('Contents', []):
                    k = obj['Key']
                    # solo archivos de evidencia de este ticket
                    fname = os.path.basename(k)
                    if not fname:
                        continue
                    # opcional: filtrar por tipos conocidos
                    if not any(fname.startswith(t + '_') for t in EVIDENCIA_TIPOS) and not fname.startswith('comb_'):
                        # incluir igual si está bajo el prefijo
                        pass
                    try:
                        data = client.get_object(Bucket=bucket, Key=k)['Body'].read()
                    except Exception:
                        continue
                    # Guardar en zip con carpeta por clave
                    arcname = f'{key}/{fname}'
                    zf.writestr(arcname, data)
                    count += 1
            except Exception as e:
                return jsonify({'error': f'Error B2: {e}'}), 500
        else:
            folder = evidencia_folder(pid, key)
            if os.path.isdir(folder):
                for fp in glob.glob(os.path.join(folder, '*.*')):
                    if not os.path.isfile(fp):
                        continue
                    fname = os.path.basename(fp)
                    try:
                        with open(fp, 'rb') as f:
                            data = f.read()
                    except Exception:
                        continue
                    zf.writestr(f'{key}/{fname}', data)
                    count += 1
    if count == 0:
        return jsonify({'error': 'No hay fotos cargadas para este ticket'}), 404
    zip_buf.seek(0)
    return Response(
        zip_buf.getvalue(),
        mimetype='application/zip',
        headers={'Content-Disposition': f'attachment; filename="evidencia_{secure_filename(key)}.zip"'}
    )


@bp.route('/api/evidencia/reporte_config')
@login_required
def api_evidencia_reporte_config():
    """Devuelve la configuración del reporte fotográfico leída de la plantilla."""
    try:
        slots = _pext_config()
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    return jsonify({'slots': slots, 'max': len(slots)})


@bp.route('/api/evidencia/reporte_xlsx/<int:pid>/<path:key>')
@login_required
def api_evidencia_reporte_xlsx(pid, key):
    """Genera el reporte fotográfico de PEXT (plantilla del cliente) con las
    fotos incrustadas en su cuadro, CM, contratista y observaciones."""
    if pid != session.get('current_proyecto_id'):
        return jsonify({'error': 'Acceso denegado'}), 403
    key = (key or '').strip()
    if not key or '/' in key or '\\' in key or '..' in key:
        return jsonify({'error': 'Clave inválida'}), 400
    if not os.path.isfile(PEX_REPORTE_XLSX):
        return jsonify({'error': 'Plantilla del reporte no disponible en el servidor.'}), 500

    record = NucleusData.query.filter_by(proyecto_id=pid, key_value=key).first()
    if not record:
        return jsonify({'error': 'Registro no encontrado'}), 404
    try:
        d = json.loads(record.data_json)
    except Exception:
        d = {}

    # Heal: si el registro aún viene del formato anterior (26), se reubica al formato
    # actual (28) y se persiste para que la interfaz también lo muestre ordenado.
    if _evidencia_migrar_legacy(d):
        record.data_json = json.dumps(d, ensure_ascii=False)
        db.session.commit()

    def _arr(campo):
        try:
            v = json.loads(d.get(campo) or '[]')
            return v if isinstance(v, list) else []
        except Exception:
            return []

    fotos = _arr('_EVIDENCIA_FOTOS')
    obs = _arr('_EVIDENCIA_OBS')
    aplica = _arr('_EVIDENCIA_APLICA')

    try:
        from openpyxl import load_workbook
        from openpyxl.drawing.image import Image as XLImage
        wb = load_workbook(PEX_REPORTE_XLSX)
    except Exception as e:
        return jsonify({'error': f'No se pudo cargar la plantilla: {e}'}), 500
    ws = wb.worksheets[0]

    # Cabecera del reporte: CM = número de WO (B5) y CONTRATISTA siempre "COBRA".
    ws['B5'] = key
    ws['J4'] = 'COBRA'
    ws['K5'] = 'COBRA'

    n_fotos = 0
    try:
        conf_slots = _pext_config() or []
    except Exception:
        conf_slots = []
    for i, slot in enumerate(conf_slots):
        top = slot.get('top')
        bottom = slot.get('bottom')
        col0 = slot.get('col0')
        col1 = slot.get('col1')
        anchor = slot.get('anchor')
        if not (top and bottom and col0 and col1):
            continue
        url = fotos[i] if i < len(fotos) else ''
        no_aplica = (i < len(aplica)) and not aplica[i]
        if url and not no_aplica:
            fname = os.path.basename(str(url).split('?')[0])
            data = _evidencia_leer(pid, key, fname)
            if data:
                box_w, box_h = _box_px(ws, top, bottom, col0, col1)
                img_w, img_h = _encajar_foto_cover(data, box_w, box_h)
                xl = XLImage(io.BytesIO(data))
                xl.width = img_w
                xl.height = img_h
                off_x = max(0, (box_w - img_w) // 2)
                off_y = max(0, (box_h - img_h) // 2)
                from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
                from openpyxl.drawing.xdr import XDRPositiveSize2D
                from openpyxl.utils.units import pixels_to_EMU
                mk = AnchorMarker(col=col0 - 1, colOff=pixels_to_EMU(off_x),
                                  row=top - 1, rowOff=pixels_to_EMU(off_y))
                xl.anchor = OneCellAnchor(_from=mk,
                                          ext=XDRPositiveSize2D(cx=pixels_to_EMU(img_w),
                                                                cy=pixels_to_EMU(img_h)))
                ws.add_image(xl)
                ws[anchor] = None  # quitar "N/A" del interior
                n_fotos += 1
        elif no_aplica:
            _n_a_en_box(ws, top, bottom, col0, col1)

        # Observaciones: la fila justo debajo del cuadro de la foto.
        # Si el gestor no escribió nada se conserva el texto por defecto que
        # viene en la plantilla; si escribió, se pisa con su observación.
        orng = None
        for mr in ws.merged_cells.ranges:
            if (mr.min_row == bottom + 1
                    and mr.min_col <= col1 and mr.max_col >= col0):
                orng = mr
                break
        o = obs[i] if i < len(obs) else ''
        o = ('' if o is None else str(o)).strip()
        if o:
            if orng:
                # Forzar que la observación quede en una sola fila.
                if orng.min_row != orng.max_row:
                    ws.unmerge_cells(str(orng))
                    ws.merge_cells(
                        start_row=orng.min_row, start_column=orng.min_col,
                        end_row=orng.min_row, end_column=orng.max_col,
                    )
                ws.cell(row=orng.min_row, column=orng.min_col).value = 'OBSERVACIONES: ' + o
            else:
                ws.cell(row=bottom + 1, column=col0).value = 'OBSERVACIONES: ' + o

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="reporte_fotografico_{secure_filename(key)}.xlsx"'}
    )