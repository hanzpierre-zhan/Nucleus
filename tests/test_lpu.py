import io,json
import pytest
from openpyxl import Workbook,load_workbook
from db import db
from models import LpuRegistro,TokenStore

@pytest.fixture(autouse=True)
def limpiar(app):
    yield
    with app.app_context():LpuRegistro.query.filter(LpuRegistro.code.like('QA-LPU-%')).delete(synchronize_session=False);db.session.commit()

def archivo(rows):
    w=Workbook();s=w.active
    for row in rows:s.append(row)
    b=io.BytesIO();w.save(b);b.seek(0);return b


def test_carga_inicial_y_codigo_default(auth_client,app):
    r=auth_client.get('/liquidaciones/api/lpu');assert r.status_code==200
    values=r.json['registros'];assert len(values)>=214
    default=next(v for v in values if v['code']=='8828126757')
    assert default['unit']=='EACH' and default['activity']==default['description_es']
    assert default['description_en']
    with app.app_context():
        from services.lpu import seed
        before=LpuRegistro.query.count();seed();assert LpuRegistro.query.count()==before


def test_crud_del_catalogo_y_duplicados(auth_client):
    payload={'code':'QA-LPU-1','unit':'UND','description_es':'Prueba','description_en':'Test','activity':'Actividad'}
    r=auth_client.post('/liquidaciones/api/lpu',json=payload);assert r.status_code==201,r.json
    rid=r.json['registro']['id']
    assert auth_client.post('/liquidaciones/api/lpu',json=payload).status_code==409
    r=auth_client.patch('/liquidaciones/api/lpu/'+str(rid),json={'activity':'Actividad corregida'})
    assert r.status_code==200 and r.json['registro']['activity']=='Actividad corregida'
    assert auth_client.patch('/liquidaciones/api/lpu/'+str(rid),json={'campo_ajeno':'X'}).status_code==400
    assert auth_client.delete('/liquidaciones/api/lpu/'+str(rid)).status_code==200
    assert auth_client.patch('/liquidaciones/api/lpu/'+str(rid),json={'unit':'m'}).status_code==404
    with pytest.raises(StopIteration):next(v for v in auth_client.get('/liquidaciones/api/lpu').json['registros'] if v['code']=='QA-LPU-1')


def test_plantilla_importacion_upsert_y_exportacion(auth_client):
    template=auth_client.get('/liquidaciones/api/lpu/plantilla');assert template.status_code==200
    w=load_workbook(io.BytesIO(template.data));header=list(next(w.active.values));assert header[0]=='Código único' and 'Unidad' in header
    upload=archivo([header,['QA-LPU-2','m','Cable','Cable','Instalación']])
    r=auth_client.post('/liquidaciones/api/lpu/importar',data={'archivo':(upload,'plantilla.xlsx')})
    assert r.status_code==200 and r.json['creados']==1
    upload=archivo([['Código único','Descripción de la partida (Esp)','LPU description (Eng)','EACH'],['QA-LPU-2','Nuevo','New','DAY']])
    r=auth_client.post('/liquidaciones/api/lpu/importar',data={'archivo':(upload,'lpu.xlsx')})
    assert r.status_code==200 and r.json['actualizados']==1
    row=next(v for v in auth_client.get('/liquidaciones/api/lpu').json['registros'] if v['code']=='QA-LPU-2')
    assert row['unit']=='DAY' and row['activity']=='Nuevo'
    export=auth_client.get('/liquidaciones/api/lpu/exportar');assert export.status_code==200
    values=list(load_workbook(io.BytesIO(export.data),data_only=True).active.values)
    assert any(v[0]=='QA-LPU-2' for v in values[1:])


def test_importacion_invalida_no_aplica_cambios_parciales(auth_client,app):
    headers=['Código único','Unidad','Descripción de la partida (Esp)','LPU description (Eng)']
    for rows in ([['QA-LPU-3','m','Correcta','Good'],['','m','Sin código','']], [['QA-LPU-3','m','Primero',''],['QA-LPU-3','m','Duplicado','']]):
        r=auth_client.post('/liquidaciones/api/lpu/importar',data={'archivo':(archivo([headers,*rows]),'lpu.xlsx')})
        assert r.status_code==400,r.json
    with app.app_context():assert not LpuRegistro.query.filter_by(code='QA-LPU-3').first()


def test_catalogo_requiere_sesion(client):
    for url,method in [('/liquidaciones/api/lpu','get'),('/liquidaciones/api/lpu','post'),('/liquidaciones/api/lpu/1','patch'),('/liquidaciones/api/lpu/1','delete')]:
        r=getattr(client,method)(url);assert r.status_code in (302,401)
