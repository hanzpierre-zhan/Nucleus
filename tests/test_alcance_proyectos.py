import io
import json

import pytest
from flask import session
from openpyxl import load_workbook
from sqlalchemy import select

from db import db
from models import AccesoProyecto, HistorialCambios, NucleusData, Proyecto, RefacturableDetalle, Usuario
from services.alcance_proyectos import limitar_opciones


@pytest.fixture
def universo(app):
    with app.app_context():
        ids = {p.nombre: p.id for p in Proyecto.query.all()}
        u = Usuario(username='alcance-prueba', rol='supervisor', password_hash='no-login')
        db.session.add(u); db.session.flush(); uid = u.id
        filas = []
        for modulo in ('Dataper', 'Material', 'SITE', 'Generadores', 'Cotizaciones', 'Rendicion', 'Combustible'):
            for familia in ('ENTEL', 'CLARO', 'INTEGRATEL', 'PEXT'):
                campo = 'CLIENTE' if modulo == 'Cotizaciones' else 'Nombre del proyecto' if modulo == 'Rendicion' else 'PROYECTO'
                d = {campo: 'Entel - Pext' if familia == 'PEXT' else familia, 'ESTADO': 'PENDIENTE',
                     'DOCUMENTO': 'SCOPE-' + familia, 'TECNICO': 'SCOPE-' + familia}
                f = NucleusData(proyecto_id=ids[modulo], key_value='SCOPE-' + modulo + '-' + familia, data_json=json.dumps(d))
                db.session.add(f); filas.append(f)
            db.session.add(AccesoProyecto(usuario_id=uid, proyecto_id=ids[modulo], restricciones='{}'))
        db.session.commit(); registros = [f.id for f in filas]
    yield uid, ids, registros
    with app.app_context():
        HistorialCambios.query.filter_by(usuario_id=uid).delete()
        RefacturableDetalle.query.filter(RefacturableDetalle.origen_id.in_(registros)).delete(synchronize_session=False)
        NucleusData.query.filter(NucleusData.id.in_(registros)).delete(synchronize_session=False)
        AccesoProyecto.query.filter_by(usuario_id=uid).delete()
        db.session.delete(db.session.get(Usuario, uid)); db.session.commit()


def asignar(app, universo, nombres):
    uid, ids, _ = universo
    with app.app_context():
        operaciones = [ids[n] for n in ('FLM - ENTEL', 'FLM - CLARO', 'FLM - INTEGRATEL', 'FLM - PEXT')]
        AccesoProyecto.query.filter(AccesoProyecto.usuario_id == uid, AccesoProyecto.proyecto_id.in_(operaciones)).delete(synchronize_session=False)
        for n in nombres:
            db.session.add(AccesoProyecto(usuario_id=uid, proyecto_id=ids[n], restricciones='{}'))
        db.session.commit()


@pytest.mark.parametrize('familias', [('ENTEL',), ('CLARO', 'INTEGRATEL'), ('ENTEL', 'CLARO', 'INTEGRATEL', 'PEXT'), ()])
def test_consultas_y_conteos_consolidan_solo_asignados(app, universo, familias):
    asignar(app, universo, ['FLM - ' + n for n in familias])
    uid, ids, _ = universo
    with app.test_request_context('/'):
        session.update(user_id=uid, rol='supervisor')
        for modulo in ('Dataper', 'Material', 'SITE', 'Generadores', 'Cotizaciones', 'Rendicion', 'Combustible'):
            q = NucleusData.query.filter(NucleusData.proyecto_id == ids[modulo], NucleusData.key_value.like('SCOPE-%'))
            assert q.count() == len(familias)
            assert {r.key_value.rsplit('-', 1)[-1] for r in q} == set(familias)
        opciones = limitar_opciones([{'nombre': 'CLIENTE', 'opciones': ['ENTEL', 'CLARO', 'INTEGRATEL', 'PEXT']}])
        assert set(opciones[0]['opciones']) == set(familias)


def test_revocacion_se_aplica_en_la_siguiente_solicitud(app, universo):
    for nombres in [('FLM - ENTEL', 'FLM - CLARO'), ('FLM - CLARO',), ()]:
        asignar(app, universo, nombres)
        with app.test_request_context('/'):
            session.update(user_id=universo[0], rol='supervisor')
            assert NucleusData.query.filter(NucleusData.proyecto_id == universo[1]['Cotizaciones'], NucleusData.key_value.like('SCOPE-%')).count() == len(nombres)


def test_exportacion_y_urls_directas_no_revelan_otro_proyecto(app, client, universo):
    asignar(app, universo, ['FLM - ENTEL'])
    uid, ids, _ = universo
    with client.session_transaction() as s:
        s.update(user_id=uid, rol='supervisor', username='alcance-prueba', current_proyecto_id=ids['Cotizaciones'], current_proyecto_nombre='Cotizaciones')
    response = client.get('/api/import/manual_template')
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data), read_only=True)
    assert 'SCOPE-Cotizaciones-CLARO' not in str(list(wb.active.values))
    wb.close()
    pagina = client.get('/')
    assert pagina.status_code == 200
    assert 'SCOPE-Cotizaciones-ENTEL' in pagina.text and 'SCOPE-Cotizaciones-CLARO' not in pagina.text
    with client.session_transaction() as s:
        s.update(current_proyecto_id=ids['Cotizaciones'], current_proyecto_nombre='Cotizaciones')
    assert client.post('/api/rows/update', json={'key': 'SCOPE-Cotizaciones-CLARO', 'field': 'MOTIVO', 'value': 'cambio'}).status_code == 404
    assert client.get('/api/evidencia/foto/%s/SCOPE-Cotizaciones-CLARO/foto.jpg' % ids['Cotizaciones']).status_code == 404
    assert client.get('/api/cotizacion/correo/%s/SCOPE-Cotizaciones-CLARO/correo.msg' % ids['Cotizaciones']).status_code == 404
    with app.app_context():
        fila = NucleusData.query.filter_by(proyecto_id=ids['Cotizaciones'], key_value='SCOPE-Cotizaciones-CLARO').one()
        assert 'MOTIVO' not in json.loads(fila.data_json)


def test_historial_tambien_respeta_proyecto(app, universo):
    asignar(app, universo, ['FLM - ENTEL'])
    uid, ids, _ = universo
    with app.app_context():
        for f in ('ENTEL', 'CLARO'):
            db.session.add(HistorialCambios(usuario_id=uid, username='alcance-prueba', proyecto_id=ids['Cotizaciones'], key_value='SCOPE-Cotizaciones-' + f, campo_modificado='MOTIVO', valor_nuevo=f))
        db.session.commit()
    with app.test_request_context('/'):
        session.update(user_id=uid, rol='supervisor')
        assert [h.valor_nuevo for h in HistorialCambios.query.filter_by(usuario_id=uid)] == ['ENTEL']


def test_liquidaciones_con_join_respeta_el_origen_sin_confundir_aliases(app, universo):
    asignar(app, universo, ['FLM - ENTEL'])
    with app.app_context():
        for f in ('ENTEL', 'CLARO'):
            origen = NucleusData.query.filter_by(proyecto_id=universo[1]['Cotizaciones'], key_value='SCOPE-Cotizaciones-' + f).one()
            db.session.add(RefacturableDetalle(origen_id=origen.id, data_json='{}'))
        db.session.commit()
    with app.test_request_context('/'):
        session.update(user_id=universo[0], rol='supervisor')
        filas = db.session.query(NucleusData, RefacturableDetalle).join(RefacturableDetalle, RefacturableDetalle.origen_id == NucleusData.id).all()
        assert len(filas) == 1 and filas[0][0].key_value == 'SCOPE-Cotizaciones-ENTEL'


def test_consultas_bulk_no_borran_filas_ajenas(app, universo):
    asignar(app, universo, ['FLM - ENTEL'])
    with app.test_request_context('/'):
        session.update(user_id=universo[0], rol='supervisor')
        afectados = NucleusData.query.filter(NucleusData.proyecto_id == universo[1]['Material'], NucleusData.key_value == 'SCOPE-Material-CLARO').delete(synchronize_session=False)
        assert afectados in (0, -1)  # SQLite no informa rowcount con una CTE.
        t = NucleusData.__table__
        assert db.session.execute(select(t.c.id).where(t.c.proyecto_id == universo[1]['Material'], t.c.key_value == 'SCOPE-Material-CLARO')).first()
        db.session.rollback()


def test_historicos_se_atribuyen_por_wo_site_y_qr_sin_adivinar(app, universo):
    asignar(app, universo, ['FLM - ENTEL'])
    uid, ids, registros = universo
    with app.app_context():
        datos = [
            ('FLM - ENTEL', 'SCOPE-WO-EN', {'Nombre de Site': 'AA001_SITE_ENTEL'}),
            ('Combustible', 'SCOPE-MOV-EN', {'WO NUMBER': 'SCOPE-WO-EN', 'QR ASIGNADO': 'QR-SCOPE'}),
            ('Combustible', 'SCOPE-INGRESO-EN', {'QR ASIGNADO': 'QR-SCOPE'}),
            ('Generadores', 'QR-SCOPE', {'QR ASIGNADO': 'QR-SCOPE'}),
            ('SITE', 'AA001', {'Código': 'AA001', 'Nombre': 'SITE_ENTEL'}),
            ('Site Name', 'SCOPE-SITE-EN', {'NOMBRE DE SITE': 'AA001_SITE_ENTEL'}),
        ]
        for modulo, key, d in datos:
            f = NucleusData(proyecto_id=ids[modulo], key_value=key, data_json=json.dumps(d))
            db.session.add(f); db.session.flush(); registros.append(f.id)
        db.session.commit()
    with app.test_request_context('/'):
        session.update(user_id=uid, rol='supervisor')
        for modulo, key, _ in datos:
            assert NucleusData.query.filter_by(proyecto_id=ids[modulo], key_value=key).one()
        # Un ingreso nuevo puede conservar el vínculo al generador sin WO.
        nuevo = NucleusData(proyecto_id=ids['Combustible'], key_value='SCOPE-NUEVO', data_json=json.dumps({'QR ASIGNADO': 'QR-SCOPE'}))
        db.session.add(nuevo); db.session.flush(); db.session.rollback()
    with app.app_context():
        ambiguo = NucleusData(proyecto_id=ids['FLM - CLARO'], key_value='SCOPE-WO-EN', data_json='{}')
        db.session.add(ambiguo); db.session.flush(); registros.append(ambiguo.id); db.session.commit()
    with app.test_request_context('/'):
        session.update(user_id=uid, rol='supervisor')
        assert NucleusData.query.filter_by(proyecto_id=ids['Generadores'], key_value='QR-SCOPE').first() is None
        assert NucleusData.query.filter_by(proyecto_id=ids['Combustible'], key_value='SCOPE-INGRESO-EN').first() is None


def test_no_reasignar_registro_a_un_proyecto_ajeno(app, universo):
    from werkzeug.exceptions import Forbidden
    asignar(app, universo, ['FLM - ENTEL'])
    with app.test_request_context('/'):
        session.update(user_id=universo[0], rol='supervisor')
        fila = NucleusData.query.filter_by(proyecto_id=universo[1]['Cotizaciones'], key_value='SCOPE-Cotizaciones-ENTEL').one()
        fila.data_json = json.dumps({'CLIENTE': 'CLARO'})
        with pytest.raises(Forbidden):
            db.session.flush()
        db.session.rollback()


def test_cache_de_opciones_no_se_comparte_con_otra_cuenta(app, client, auth_client, universo):
    asignar(app, universo, ['FLM - ENTEL'])
    uid, ids, registros = universo
    with app.app_context():
        for f in ('ENTEL', 'CLARO'):
            row = NucleusData(proyecto_id=ids['SITE'], key_value='SCOPE-CACHE-' + f,
                              data_json=json.dumps({'PROYECTO': f, 'Nombre': 'SCOPE-CACHE-' + f}))
            db.session.add(row); db.session.flush(); registros.append(row.id)
        db.session.commit()
    assert 'SCOPE-CACHE-CLARO' in auth_client.get('/api/detalle/opciones').text
    with client.session_transaction() as s:
        s.update(user_id=uid, rol='supervisor', current_proyecto_id=ids['FLM - ENTEL'], current_proyecto_nombre='FLM - ENTEL')
    respuesta = client.get('/api/detalle/opciones')
    assert respuesta.status_code == 200
    assert 'SCOPE-CACHE-ENTEL' in respuesta.text and 'SCOPE-CACHE-CLARO' not in respuesta.text


@pytest.mark.parametrize('rol', ['gestor', 'contrata'])
@pytest.mark.parametrize('metodo,ruta', [
    ('get', '/api/import/manual_template'),
    ('post', '/api/import/preview'),
    ('post', '/api/import/process'),
    ('get', '/api/cotizacion/items_template'),
    ('post', '/api/cotizacion/items_import'),
    ('get', '/api/master/template/materiales'),
    ('post', '/api/master/bulk_import/materiales'),
    ('post', '/api/rows/delete'),
    ('post', '/api/cotizacion/eliminar'),
    ('post', '/api/evidencia/eliminar'),
])
def test_roles_operativos_no_importan_ni_borran(app, client, universo, rol, metodo, ruta):
    uid, ids, _ = universo
    with app.app_context():
        db.session.get(Usuario, uid).rol = rol
        db.session.commit()
    with client.session_transaction() as s:
        # Incluso una cookie con el rol anterior no conserva permisos.
        s.update(user_id=uid, rol='zeno', current_proyecto_id=ids['SITE'], current_proyecto_nombre='SITE')
    assert getattr(client, metodo)(ruta, json={}).status_code == 403


def test_soporte_solo_proyecto_asignado(app, client, universo):
    from services.utilidades import get_menu_proyectos
    uid, ids, _ = universo
    asignar(app, universo, ['FLM - CLARO'])
    with app.app_context():
        db.session.get(Usuario, uid).rol = 'suport'
        db.session.commit()
    with app.test_request_context('/'):
        session.update(user_id=uid, rol='suport')
        assert ids['FLM - ENTEL'] not in {p.id for p in get_menu_proyectos(uid, 'suport')}
        filas = NucleusData.query.filter_by(proyecto_id=ids['SITE']).all()
        assert {r.key_value for r in filas if r.key_value.startswith('SCOPE-')} == {'SCOPE-SITE-CLARO'}
    with client.session_transaction() as s:
        s.update(user_id=uid, rol='suport', current_proyecto_id=ids['FLM - CLARO'], current_proyecto_nombre='FLM - CLARO')
    client.get('/switch_project/%s' % ids['FLM - ENTEL'])
    with client.session_transaction() as s:
        assert s['current_proyecto_id'] != ids['FLM - ENTEL']
    with client.session_transaction() as s:
        s.update(current_proyecto_id=ids['FLM - ENTEL'], current_proyecto_nombre='FLM - ENTEL')
    assert client.get('/api/import/manual_template').status_code == 403


def test_supervisor_descarga_importa_y_elimina_solo_su_proyecto(app, client, universo):
    uid, ids, registros = universo
    asignar(app, universo, ['FLM - CLARO'])
    with client.session_transaction() as s:
        s.update(user_id=uid, rol='supervisor', current_proyecto_id=ids['SITE'], current_proyecto_nombre='SITE')
    pagina = client.get('/')
    assert pagina.status_code == 200
    assert 'id="btn-top-import"' in pagina.text
    assert 'id="btn-top-plantilla"' in pagina.text
    assert 'id="btn-delete-rows"' in pagina.text
    plantilla = client.get('/api/import/manual_template')
    assert plantilla.status_code == 200
    libro = load_workbook(io.BytesIO(plantilla.data))
    assert 'SCOPE-SITE-ENTEL' not in str(list(libro.active.values))
    libro.close()
    contenido = 'Código,Nombre,PROYECTO\nSCOPE-SUPERVISOR,Nuevo,CLARO\n'.encode('utf-8')
    respuesta = client.post('/api/import/process', data={
        'type': 'manual_cols', 'file': (io.BytesIO(contenido), 'sites.csv')})
    assert respuesta.status_code == 200, respuesta.json
    assert respuesta.json['added'] == 1
    with app.app_context():
        fila = NucleusData.query.filter_by(proyecto_id=ids['SITE'], key_value='SCOPE-SUPERVISOR').one()
        registros.append(fila.id)
    fuera = contenido.replace(b'CLARO', b'ENTEL')
    respuesta = client.post('/api/import/process', data={
        'type': 'manual_cols', 'file': (io.BytesIO(fuera), 'sites.csv')})
    assert respuesta.status_code == 403
    respuesta = client.post('/api/rows/delete', json={'keys':['SCOPE-SUPERVISOR', 'SCOPE-SITE-ENTEL']})
    assert respuesta.status_code == 200, respuesta.json
    with app.app_context():
        assert NucleusData.query.filter_by(proyecto_id=ids['SITE'], key_value='SCOPE-SUPERVISOR').first() is None
        assert NucleusData.query.filter_by(proyecto_id=ids['SITE'], key_value='SCOPE-SITE-ENTEL').first() is not None
