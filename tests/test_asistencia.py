import json
import pytest
from db import db
from models import (Proyecto,NucleusData,Usuario,AccesoProyecto,AsistenciaDiaria,AsistenciaCambio,CuentaEspacio,AccesoEspacio,PerfilEspacio)


@pytest.fixture(params=[False,True])
def equipo(app,client,monkeypatch,request):
    monkeypatch.setitem(app.config,'ESPACIOS_HABILITADOS',True)
    with app.app_context():
        ids={p.nombre:p.id for p in Proyecto.query.all()}
        u=Usuario(username='asistencia-gestora',nombre='Gestora de prueba',rol='gestor',password_hash='sin-login');db.session.add(u);db.session.flush();uid=u.id
        perfil=None
        if request.param:
            perfil=PerfilEspacio(nombre='Asistencia prueba',permisos_json=json.dumps({'personal':['consultar','editar','exportar']}));db.session.add(perfil);db.session.flush()
            db.session.add(CuentaEspacio(usuario_id=uid))
            for n in ('FLM - CLARO','FLM - INTEGRATEL'):db.session.add(AccesoEspacio(usuario_id=uid,proyecto_id=ids[n],perfil_id=perfil.id))
        else:
            for n in ('FLM - CLARO','FLM - INTEGRATEL','Dataper'):db.session.add(AccesoProyecto(usuario_id=uid,proyecto_id=ids[n],restricciones='{}'))
        personas=[]
        for n in ('FLM - CLARO','FLM - INTEGRATEL','FLM - ENTEL'):
            p=NucleusData(proyecto_id=ids['Dataper'],key_value='ASIS-'+n,data_json=json.dumps({'PROYECTO':n,'TECNICO':'Persona '+n,'ESTADO':'ACTIVO'}));db.session.add(p);personas.append(p)
        db.session.commit();pids=[p.id for p in personas];perfil_id=perfil.id if perfil else None
    with client.session_transaction() as s:s.update(user_id=uid,username='asistencia-gestora',rol='gestor')
    assert client.get('/personal/asistencia').status_code==200
    with client.session_transaction() as s:token=s['asistencia_csrf']
    yield {'c':client,'ids':ids,'uid':uid,'personas':pids,'headers':{'X-Asistencia-CSRF':token}}
    with app.app_context():
        marcas=[r.id for r in AsistenciaDiaria.query.filter(AsistenciaDiaria.personal_id.in_(pids))]
        AsistenciaCambio.query.filter(AsistenciaCambio.asistencia_id.in_(marcas)).delete(synchronize_session=False)
        AsistenciaDiaria.query.filter(AsistenciaDiaria.personal_id.in_(pids)).delete(synchronize_session=False)
        AccesoEspacio.query.filter_by(usuario_id=uid).delete();CuentaEspacio.query.filter_by(usuario_id=uid).delete()
        if perfil_id:PerfilEspacio.query.filter_by(id=perfil_id).delete()
        AccesoProyecto.query.filter_by(usuario_id=uid).delete();NucleusData.query.filter(NucleusData.id.in_(pids)).delete(synchronize_session=False);Usuario.query.filter_by(id=uid).delete();db.session.commit()


def test_asistencia_solo_proyectos_autorizados_sin_faltas_inferidas(equipo):
    c=equipo['c'];j=c.get('/api/asistencia/dia?fecha=2026-10-17').json
    assert {r['proyecto'] for r in j['registros']}=={'FLM - CLARO','FLM - INTEGRATEL'}
    assert all(r['estado']=='PENDIENTE' and r['version']==0 for r in j['registros'])
    assert c.get('/api/asistencia/dia?fecha=2026-10-17&proyecto='+str(equipo['ids']['FLM - ENTEL'])).status_code==403
    assert c.get('/api/asistencia/mensual?mes=2026-10').json['marcas']==0
    plan=c.get('/api/asistencia/planilla?mes=2026-02').json
    assert plan['dias']==28
    assert {r['proyecto'] for r in plan['registros']}=={'FLM - CLARO','FLM - INTEGRATEL'}
    assert all(r['marcas']=={} for r in plan['registros'])
    assert c.get('/api/asistencia/planilla?mes=2026-13').status_code==400
    assert c.get('/api/asistencia/planilla?mes=2026-10&proyecto='+str(equipo['ids']['FLM - ENTEL'])).status_code==403


def test_guardado_historial_y_control_de_version(app,equipo):
    c=equipo['c'];r=c.get('/api/asistencia/dia?fecha=2026-10-17').json['registros'][0]
    marca={k:r[k] for k in ('personal_id','proyecto_id','version')};marca.update(estado='PRESENTE',observacion='Equipo verificado')
    payload={'fecha':'2026-10-17','registros':[marca]}
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json=payload).status_code==200
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json=payload).status_code==409
    marca.update(version=1,estado='PERMISO',observacion='Permiso registrado')
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json=payload).status_code==200
    diario=c.get('/api/asistencia/dia?fecha=2026-10-17').json['registros']
    assert len([x for x in diario if x['estado']=='PERMISO'])==1
    assert len([x for x in diario if x['estado']=='PENDIENTE'])==1
    with app.app_context():
        row=AsistenciaDiaria.query.filter_by(personal_id=marca['personal_id']).one();assert row.version==2
        assert AsistenciaCambio.query.filter_by(asistencia_id=row.id).count()==2
    mensual=c.get('/api/asistencia/mensual?mes=2026-10').json
    assert mensual['marcas']==1 and mensual['resumen'][0]['estados']['PERMISO']==1
    plan=c.get('/api/asistencia/planilla?mes=2026-10').json
    persona=next(p for p in plan['registros'] if p['personal_id']==marca['personal_id'])
    assert persona['marcas']['17']['estado']=='PERMISO' and persona['marcas']['17']['version']==2
    csv=c.get('/api/asistencia/mensual?mes=2026-10&formato=csv');assert csv.status_code==200 and 'text/csv' in csv.content_type and 'Permiso registrado' in csv.text


def test_historial_conserva_proyecto_al_cesar_y_csv_seguro(app,equipo):
    c=equipo['c'];r=c.get('/api/asistencia/dia?fecha=2026-10-17').json['registros'][0]
    marca={k:r[k] for k in ('personal_id','proyecto_id','version')}
    marca.update(estado='PRESENTE',observacion='=1+1')
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json={'fecha':'2026-10-17','registros':[marca]}).status_code==200
    with app.app_context():
        p=db.session.get(NucleusData,r['personal_id']);d=json.loads(p.data_json)
        d.update(ESTADO='CESADO',PROYECTO='FLM - ENTEL');p.data_json=json.dumps(d);db.session.commit()
    historico=c.get('/api/asistencia/dia?fecha=2026-10-17').json['registros']
    assert any(x['personal_id']==r['personal_id'] and x['proyecto_id']==r['proyecto_id'] for x in historico)
    assert all(x['personal_id']!=r['personal_id'] for x in c.get('/api/asistencia/planilla?mes=2026-10').json['registros'])
    cesados=c.get('/api/asistencia/planilla?mes=2026-10&personal=cesados').json['registros']
    assert len(cesados)==1 and cesados[0]['personal_id']==r['personal_id']
    assert any(x['personal_id']==r['personal_id'] for x in c.get('/api/asistencia/planilla?mes=2026-10&personal=ambos').json['registros'])
    assert c.get('/api/asistencia/planilla?mes=2026-10&personal=otro').status_code==400
    marca.update(version=1,estado='PERMISO')
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json={'fecha':'2026-10-17','registros':[marca]}).status_code==200
    assert "'=1+1" in c.get('/api/asistencia/mensual?mes=2026-10&formato=csv').text
    import io
    from openpyxl import load_workbook
    excel=c.get('/api/asistencia/mensual?mes=2026-10&personal=cesados&formato=xlsx')
    assert excel.status_code==200 and '.xlsx' in excel.headers['Content-Disposition']
    hoja=load_workbook(io.BytesIO(excel.data)).active
    assert hoja.max_row==2 and hoja['F2'].value=='=1+1' and hoja['F2'].data_type=='s'
    assert hoja['A2'].value.day==17 and hoja.freeze_panes=='D2'
    vacio=c.get('/api/asistencia/mensual?mes=2026-10&personal=activos&formato=xlsx')
    assert load_workbook(io.BytesIO(vacio.data)).active.max_row==1
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json={'fecha':'2026-10-18','registros':[{**marca,'version':0}]}).status_code==403


def test_batch_invalido_no_guarda_parcialmente(app,equipo):
    c=equipo['c'];r=c.get('/api/asistencia/dia?fecha=2026-10-17').json['registros'][0]
    marca={k:r[k] for k in ('personal_id','proyecto_id','version')};marca.update(estado='PRESENTE',observacion='')
    otra={**marca,'personal_id':equipo['personas'][2],'proyecto_id':equipo['ids']['FLM - ENTEL']}
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json={'fecha':'2026-10-17','registros':[marca,otra]}).status_code==403
    assert c.get('/api/asistencia/mensual?mes=2026-10').json['marcas']==0
    assert c.post('/api/asistencia/dia',json={'fecha':'2026-10-17','registros':[marca]}).status_code==403
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json={'fecha':'mal','registros':[marca]}).status_code==400


def test_contrata_no_marca_y_personal_con_asistencia_no_se_elimina(app,auth_client,equipo):
    c=equipo['c'];r=c.get('/api/asistencia/dia?fecha=2026-10-17').json['registros'][0]
    marca={k:r[k] for k in ('personal_id','proyecto_id','version')};marca.update(estado='PRESENTE',observacion='')
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json={'fecha':'2026-10-17','registros':[marca]}).status_code==200
    auth_client.get('/switch_project/%s?next=/'%equipo['ids']['Dataper'])
    with app.app_context():key=db.session.get(NucleusData,marca['personal_id']).key_value
    assert auth_client.post('/api/rows/delete',json={'keys':[key]}).status_code==409
    with app.app_context():db.session.get(Usuario,equipo['uid']).rol='contrata';db.session.commit()
    marca['version']=1;marca['estado']='FALTA'
    assert c.post('/api/asistencia/dia',headers=equipo['headers'],json={'fecha':'2026-10-17','registros':[marca]}).status_code==403
