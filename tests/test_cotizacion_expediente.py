import io,json,zipfile
from datetime import datetime
import pytest
from openpyxl import load_workbook
from PIL import Image
from db import db
from models import NucleusData
from test_sustento_por_wo import caso,foto
from services.cotizacion_expediente import generar
from services.cotizacion_ficha import items_cotizacion

@pytest.mark.parametrize('client,manager',[('CLARO','Dennis Unton'),('INTEGRATEL','Dennis Unton'),('ENTEL','Hector Pujada')])
def test_plantilla_por_cliente_y_fotos_completas_sin_formulas_rotas(client,manager,tmp_path):
    datos={'CLIENTE':client,'SUPERVISOR':'No usar este valor','N° COTIZACION':'HW-2026-000345',
           'ITEMS_JSON':json.dumps([{'texto':'Compra de cable','und':'UND','cantidad':18},{'texto':'Instalación','und':'m','cantidad':5}])}
    items=items_cotizacion(datos)
    manual={'work_responsible':'Gestor Responsable','ajb_ejb':'AJB-1','expense':'EXP-2',
            'site_name':'Sitio de prueba','site_id':'00123','site_address':'Dirección del sitio','site_type':'Outdoor',
            'district':'Distrito','province':'Provincia','region':'Región','location_picture':'ubicacion',
            'item_details':{i['key']:{'unique_code':'8828126757','unit':'EACH','description_es':'Partida ES',
                            'description_en':'LPU EN','remarks':'Actividad realizada '+str(n)} for n,i in enumerate(items,1)},
            'item_photos':{items[0]['key']:[{'url':'foto-'+str(n),'phase':'before'} for n in range(5)]+
                                         [{'url':'durante','phase':'during'},{'url':'despues','phase':'after'}]}}
    photo=io.BytesIO();Image.new('RGB',(800,400),'navy').save(photo,'PNG')
    loaded=[]
    def read(url):loaded.append(url);return photo.getvalue()
    output=generar(datos,manual,{'_BITACORA_INICIO':'Antecedentes','_BITACORA_PROCESO':'Objetivos',
        '_BITACORA_CIERRE':'Trabajo realizado','INICIO EJECUCION':'2026-10-01 08:30','FIN EJECUCION':'2026-10-02 18:00'},'WO-1',read,datetime(2026,10,23,12,0))
    book=load_workbook(output)
    assert book['Carátula']['E43'].value.date().isoformat()=='2026-10-23'
    assert book['Carátula']['E45'].value=='HW-2026-000345'
    assert book['Carátula']['E44'].value=='AJB-1'
    general=book['Información general']
    assert general['E3'].value=='COBRA PERU S.A.C.'
    assert general['E4'].value=='20253881438'
    assert general['K3'].value==general['K4'].value==manager
    assert general['K5'].value=='Gestor Responsable'
    assert general['E9'].value=='Sitio de prueba' and general['K9'].value=='00123'
    assert book['Detalles']['B3'].value=='Antecedentes'
    assert len(general._images)==1 and len(book['Carátula']._images)==2
    parts=book['Partidas']
    assert [parts.cell(r,13).value for r in (9,18)]==[1,2]
    assert [parts.cell(r,8).value for r in (10,19)]==[1,1]
    assert [b.id for b in parts.row_breaks.brk]==[17]
    assert parts.max_row<=26
    assert parts['E13'].value=='Compra de cable'
    assert parts['E16'].value=='Actividad realizada 1'
    assert loaded==['ubicacion','foto-0','foto-1','foto-2','foto-3','foto-4','durante','despues']
    assert len(parts._images)==7
    assert all(image.anchor._from.row==14 for image in parts._images)
    for sheet in book:
        for row in sheet:
            assert all(c.data_type not in ('f','e') for c in row),sheet.title
    output.seek(0);(tmp_path/('expediente-'+client+'.xlsx')).write_bytes(output.read())
    output.seek(0)
    with zipfile.ZipFile(output) as archive:
        assert 'xl/media/nucleus_header.png' in archive.namelist()
        assert b'legacyDrawingHF' in archive.read('xl/worksheets/sheet2.xml')

def test_descarga_incluye_fotos_guardadas_y_no_cambia_la_cotizacion(app,auth_client,caso,monkeypatch,tmp_path):
    import blueprints.cotizacion as cot
    monkeypatch.setattr(cot,'evidencia_usa_b2',lambda:False)
    monkeypatch.setattr(cot,'_ubicacion_folder',lambda cid:str(tmp_path/str(cid)))
    with app.app_context():
        row=db.session.get(NucleusData,caso['ids'][2]);d=json.loads(row.data_json)
        d['ITEMS_JSON']=json.dumps([{'texto':'Compra de cable','und':'UND','cantidad':5}]);d['N° COTIZACION']='QA-EXPEDIENTE'
        row.data_json=json.dumps(d);db.session.commit()
    auth_client.get('/switch_project/'+str(caso['cot']))
    item=auth_client.get('/api/cotizacion/sustento?key=COT-SUST-TEST').json['items'][0]
    upload=auth_client.post('/api/cotizacion/sustento/item-foto',data={'key':'COT-SUST-TEST','wo':'WO-SUST-1','item_key':item['key'],
        'phase':'before','foto':(foto(),'foto.png')})
    assert upload.status_code==200,upload.json
    result=auth_client.get('/api/cotizacion/sustento/expediente',query_string={'key':'COT-SUST-TEST','wo':'WO-SUST-1'})
    assert result.status_code==200,result.data[:300]
    assert 'attachment' in result.headers['Content-Disposition']
    assert len(load_workbook(io.BytesIO(result.data))['Partidas']._images)==1
    assert auth_client.get('/api/cotizacion/sustento/expediente?key=COT-SUST-TEST&wo=OTRA').status_code==400
    assert auth_client.get('/api/cotizacion/sustento/expediente?key=NO-EXISTE&wo=WO-SUST-1').status_code==404
    with app.app_context():
        row=db.session.get(NucleusData,caso['ids'][2]);d=json.loads(row.data_json)
        assert json.loads(d['ITEMS_JSON'])[0]['cantidad']==5
        d['CLIENTE']='PEXT';row.data_json=json.dumps(d);db.session.commit()
    assert auth_client.get('/api/cotizacion/sustento/expediente?key=COT-SUST-TEST&wo=WO-SUST-1').status_code==400

def test_catalogo_responsables_persistente_y_sin_cambiar_opciones_supervisor(app,auth_client,caso):
    auth_client.get('/switch_project/'+str(caso['cot']))
    result=auth_client.patch('/api/cotizacion/sustento/responsables',json={'opciones':['Responsable A','Responsable B','Responsable A']})
    assert result.status_code==200 and result.json['opciones']==['Responsable A','Responsable B']
    assert auth_client.get('/api/cotizacion/sustento/responsables').json['opciones']==result.json['opciones']
    assert auth_client.patch('/api/cotizacion/sustento/responsables',json={'opciones':['']}).status_code==400
    assert app.test_client().get('/api/cotizacion/sustento/expediente').status_code in (302,401)

@pytest.mark.parametrize('count',[1,2,3,4,5,6,13,30])
def test_fotos_independientes_ordenadas_sin_solaparse_ni_salir_del_espacio(count):
    from openpyxl import Workbook
    from services.cotizacion_expediente import fotos_independientes,width
    sheet=Workbook().active
    sheet.row_dimensions[15].height=199.5
    contents=[]
    sizes=[(900,450),(400,900),(640,640)]
    for index in range(count):
        content=io.BytesIO();Image.new('RGB',sizes[index%3],(index*7%255,40,100)).save(content,'PNG')
        contents.append(content.getvalue())
    fotos_independientes(sheet,contents,6,15,250,266)
    assert len(sheet._images)==count
    rects=[]
    for index,image in enumerate(sheet._images):
        anchor=image.anchor
        x=sum(width(sheet,c,c) for c in range(6,anchor._from.col+1))+anchor._from.colOff/9525
        y=anchor._from.rowOff/9525
        w=anchor.ext.cx/9525;h=anchor.ext.cy/9525
        assert anchor._from.col>=5 and anchor._from.row==14
        assert anchor._from.colOff/9525<width(sheet,anchor._from.col+1,anchor._from.col+1)
        assert x>=0 and y>=0 and x+w<=250 and y+h<=266
        assert abs(w/h-sizes[index%3][0]/sizes[index%3][1])<0.08
        with Image.open(io.BytesIO(image._data())) as original:assert original.size==sizes[index%3]
        rects.append((x,y,x+w,y+h))
    for i,a in enumerate(rects):
        for b in rects[i+1:]:assert a[2]<=b[0] or b[2]<=a[0] or a[3]<=b[1] or b[3]<=a[1]

@pytest.mark.parametrize('size',[(1600,900),(900,1600)])
def test_ubicacion_aprovecha_el_espacio_sin_deformar_ni_recortar(size):
    from services.cotizacion_expediente import width
    content=io.BytesIO();Image.new('RGB',size,'teal').save(content,'PNG')
    output=generar({'CLIENTE':'ENTEL','ITEMS_JSON':json.dumps([{'texto':'Trabajo'}])},
        {'location_picture':'foto'}, {},'',lambda url:content.getvalue(),datetime(2026,10,26))
    sheet=load_workbook(output)['Información general']
    image=sheet._images[0];anchor=image.anchor
    x=sum(width(sheet,c,c) for c in range(2,anchor._from.col+1))+anchor._from.colOff/9525
    y=anchor._from.rowOff/9525
    w=anchor.ext.cx/9525;h=anchor.ext.cy/9525
    available_width=width(sheet,2,13);available_height=sheet.row_dimensions[13].height*96/72
    assert x+w<=available_width and y+h<=available_height
    assert y>=28
    assert abs(w/h-size[0]/size[1])<0.01
    assert sheet.row_dimensions[13].height==253.5
    assert min(available_width-w,available_height-28-h)<=9

def test_dos_verticales_y_una_horizontal_como_el_ejemplo(tmp_path):
    from services.cotizacion_expediente import distribuir_fotos
    rectangles=distribuir_fotos(['vertical1','vertical2','horizontal'],[0.5,0.5,2.4],330,333)
    first,second,last=rectangles
    assert abs(first[1]-second[1])<1
    assert first[0]+first[2]<=second[0]
    assert last[1]>=first[1]+first[3]
    assert last[2]>first[2]+second[2]
    images={}
    for name,size,color in [('vertical1',(400,800),'teal'),('vertical2',(400,800),'navy'),('horizontal',(1200,500),'darkred')]:
        content=io.BytesIO();Image.new('RGB',size,color).save(content,'PNG');images[name]=content.getvalue()
    datos={'CLIENTE':'ENTEL','ITEMS_JSON':json.dumps([{'texto':'Fotos de orientaciones mixtas'}])}
    item=items_cotizacion(datos)[0]
    output=generar(datos,{'location_picture':'horizontal','item_photos':{item['key']:[{'url':key,'phase':'during'} for key in images]}},
        {},'',lambda url:images[url],datetime(2026,10,26))
    (tmp_path/'ejemplo-mixto.xlsx').write_bytes(output.getvalue())
