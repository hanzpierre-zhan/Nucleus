"""El reporte de sustento no debe perder las fechas de ejecución.

Se cubren los dos caminos reales: la fecha viaja en el WO y, si el WO la
perdió (sync/reimportación) o quedó vacía, se recupera de la cotización.
"""
import io
import json

from openpyxl import load_workbook

from db import db
from models import NucleusData
from test_sustento_por_wo import caso, foto  # noqa: F401  (fixtures)


def _guardar_cotizacion(app, caso, **cambios):
    with app.app_context():
        fila = db.session.get(NucleusData, caso['ids'][2])
        datos = json.loads(fila.data_json)
        datos.setdefault('ITEMS_JSON', json.dumps([{'texto': 'Trabajo', 'und': 'UND', 'cantidad': 1}]))
        datos.update(cambios)
        fila.data_json = json.dumps(datos)
        db.session.commit()


def _fechas_del_reporte(auth_client):
    respuesta = auth_client.get('/api/cotizacion/sustento/expediente',
                                query_string={'key': 'COT-SUST-TEST', 'wo': 'WO-SUST-1'})
    assert respuesta.status_code == 200, respuesta.data[:400]
    general = load_workbook(io.BytesIO(respuesta.data))['Información general']
    return general['E6'].value, general['K6'].value


def test_wo_con_fechas_las_muestra_en_el_reporte(app, auth_client, caso):
    _guardar_cotizacion(app, caso, **{'NUMERO WO': 'WO-SUST-1'})
    auth_client.get('/switch_project/%s' % caso['cot'])
    respuesta = auth_client.post('/api/cotizacion/sustento/wo', json={
        'key': 'COT-SUST-TEST', 'wo': 'WO-SUST-1',
        'evidencia': {'_BITACORA_INICIO': 'Antecedentes',
                      'INICIO EJECUCION': '2026-10-01 08:30:00',
                      'FIN EJECUCION': '2026-10-02 18:00:00'}})
    assert respuesta.status_code == 200, respuesta.json
    assert _fechas_del_reporte(auth_client) == ('2026-10-01 08:30:00', '2026-10-02 18:00:00')


def test_wo_sin_fechas_las_recupera_de_la_cotizacion(app, auth_client, caso):
    """Un WO ya migrado no debe ocultar las fechas de obra de la cotización."""
    with app.app_context():
        fila = db.session.get(NucleusData, caso['ids'][0])
        fila.data_json = json.dumps({'_BITACORA_INICIO': 'x', '_COT_EVIDENCIA_MIGRADA': '1'})
        db.session.commit()
    _guardar_cotizacion(app, caso, **{'NUMERO WO': 'WO-SUST-1',
                                      'FECHA INICIO OBRA': '2026-10-21',
                                      'FECHA FIN OBRA': '2026-10-23'})
    auth_client.get('/switch_project/%s' % caso['cot'])
    assert _fechas_del_reporte(auth_client) == ('2026-10-21 00:00:00', '2026-10-23 00:00:00')


def test_guardar_sustento_respalda_las_fechas_en_la_cotizacion(app, auth_client, caso):
    _guardar_cotizacion(app, caso, **{'NUMERO WO': 'WO-SUST-1'})
    auth_client.get('/switch_project/%s' % caso['cot'])
    respuesta = auth_client.post('/api/cotizacion/sustento/wo', json={
        'key': 'COT-SUST-TEST', 'wo': 'WO-SUST-1',
        'evidencia': {'INICIO EJECUCION': '2026-11-02 07:15:00',
                      'FIN EJECUCION': '2026-11-04 19:45:00'}})
    assert respuesta.status_code == 200, respuesta.json
    with app.app_context():
        datos = json.loads(db.session.get(NucleusData, caso['ids'][2]).data_json)
        assert datos['INICIO EJECUCION'] == '2026-11-02 07:15:00'
        assert datos['FIN EJECUCION'] == '2026-11-04 19:45:00'
        assert datos['FECHA INICIO OBRA'] == '2026-11-02'
        assert datos['FECHA FIN OBRA'] == '2026-11-04'
    assert _fechas_del_reporte(auth_client) == ('2026-11-02 07:15:00', '2026-11-04 19:45:00')


def test_generar_usa_las_fechas_de_la_cotizacion_si_la_evidencia_no_las_tiene():
    from datetime import datetime
    from services.cotizacion_expediente import generar
    from services.cotizacion_ficha import items_cotizacion

    datos = {'CLIENTE': 'ENTEL', 'N° COTIZACION': 'HW-TEST',
             'ITEMS_JSON': json.dumps([{'texto': 'Trabajo'}]),
             'FECHA INICIO OBRA': '2026-10-05', 'FECHA FIN OBRA': '2026-10-06'}
    salida = generar(datos, {'work_responsible': 'Responsable'}, {}, '', lambda url: b'',
                     datetime(2026, 10, 27, 9, 0))
    general = load_workbook(salida)['Información general']
    assert general['E6'].value == '2026-10-05 00:00:00'
    assert general['K6'].value == '2026-10-06 00:00:00'
    assert items_cotizacion(datos)[0]['texto'] == 'Trabajo'
