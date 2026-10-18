# -*- coding: utf-8 -*-
"""Equivalencia del .xlsx: pandas.to_excel (antes) vs openpyxl directo (ahora).

Compara celda a celda los dos caminos para los dos formatos del endpoint
(FLM/PEXT y genérico), incluyendo nombres de columna repetidos.
"""
import io
import json
import os
import sqlite3
import sys

BASE = r'C:\Mega\Proyect\Nucleus'
sys.path.insert(0, BASE)

import pandas as pd  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

DB = os.path.join(os.environ['TEMP'], 'bench_nucleus.db')
con = sqlite3.connect(DB)
filas_db = [(k, json.loads(dj)) for k, dj in
            con.execute('select key_value,data_json from nucleus_data where proyecto_id=14 limit 300')]

DETALLE = [
    'Fecha de creación (WO Creation date)', 'Nombre de Site', 'Autin TT',
    'Departamento', 'Fault Level', 'Estado de la tarea (WO State)',
    'Prioridad del Site', 'Provincia', 'Distrito',
]
MANUAL = ['SERVICIO', 'CIUDAD', 'Departamento', 'Provincia',  # repetidos a propósito
          'LATITUD', 'LONGITUD', 'MOTIVO DE AVERÍA', 'SOLUCIÓN',
          'TECNICO ASIGNADO', 'CONTRATA', 'SISTEMAS', 'BITACORA']


def _antes(pk_name, detalle_cols, manual_cols, parsed, flm):
    data = {}
    if flm:
        data[pk_name] = [k for k, _d in parsed]
        for col in detalle_cols:
            data[col] = [str((d.get(col) or '')) for _k, d in parsed]
        for col in manual_cols:
            data[col] = [str((d.get(col) or '')) for _k, d in parsed]
    else:
        data[pk_name] = [k for k, _d in parsed]
        for col in manual_cols:
            data[col] = [''] * len(parsed)
    df = pd.DataFrame(data)
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Plantilla')
        ws = writer.sheets['Plantilla']
        tope = 38 if flm else 40
        if flm:
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, tope)
        if flm:
            fill_det = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
            fill_ges = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            for idx, col_name in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=idx)
                cell.font = Font(bold=True, size=9)
                if col_name in detalle_cols:
                    cell.fill = fill_det
                elif col_name in manual_cols:
                    cell.fill = fill_ges
    out.seek(0)
    return out.getvalue()


def _ahora(pk_name, detalle_cols, manual_cols, parsed, flm):
    if flm:
        columnas = [pk_name] + list(dict.fromkeys(detalle_cols + manual_cols))
        filas = []
        for clave, d in parsed:
            fila = {pk_name: clave}
            for col in detalle_cols + manual_cols:
                fila[col] = str(d.get(col) or '')
            filas.append([fila[col] for col in columnas])
        tope = 38
    else:
        por_columna = {pk_name: [k for k, _d in parsed]}
        for col in manual_cols:
            por_columna[col] = [''] * len(parsed)
        columnas = list(por_columna.keys())
        filas = [[por_columna[col][i] for col in columnas] for i in range(len(parsed))]
        tope = 40

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Plantilla'
    hoja.append(list(columnas))
    anchos = [len(str(c)) for c in columnas]
    for fila in filas:
        hoja.append(fila)
        for i, valor in enumerate(fila):
            largo = len(str(valor or ''))
            if largo > anchos[i]:
                anchos[i] = largo
    for i, largo in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(i)].width = min(largo + 4, tope)
    if flm:
        hoja.freeze_panes = 'A2'
        hoja.auto_filter.ref = hoja.dimensions
        fill_det = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
        fill_ges = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        estilos = {c: fill_ges for c in manual_cols}
        estilos.update({c: fill_det for c in detalle_cols})
        fuente = Font(bold=True, size=9)
        for idx, nombre_col in enumerate(columnas, 1):
            celda = hoja.cell(row=1, column=idx)
            celda.font = fuente
            if nombre_col in estilos:
                celda.fill = estilos[nombre_col]
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def comparar(etiqueta, pk, detalle, manual, parsed, flm):
    a = openpyxl.load_workbook(io.BytesIO(_antes(pk, detalle, manual, parsed, flm)))['Plantilla']
    b = openpyxl.load_workbook(io.BytesIO(_ahora(pk, detalle, manual, parsed, flm)))['Plantilla']
    dif = []
    if a.dimensions != b.dimensions:
        dif.append(f'dimensiones {a.dimensions} != {b.dimensions}')
    if a.freeze_panes != b.freeze_panes:
        dif.append(f'freeze {a.freeze_panes} != {b.freeze_panes}')
    if a.auto_filter.ref != b.auto_filter.ref:
        dif.append(f'autofiltro {a.auto_filter.ref} != {b.auto_filter.ref}')
    for f in range(1, a.max_row + 1):
        for c in range(1, a.max_column + 1):
            if a.cell(row=f, column=c).value != b.cell(row=f, column=c).value:
                dif.append(f'valor {f},{c}: {a.cell(row=f,column=c).value!r} != {b.cell(row=f,column=c).value!r}')
    for c in range(1, a.max_column + 1):
        letra = get_column_letter(c)
        if a.column_dimensions[letra].width != b.column_dimensions[letra].width:
            dif.append(f'ancho {letra}: {a.column_dimensions[letra].width} != {b.column_dimensions[letra].width}')
    for c in range(1, a.max_column + 1):
        ca, cb = a.cell(row=1, column=c), b.cell(row=1, column=c)
        if (ca.font.bold, ca.font.size) != (cb.font.bold, cb.font.size):
            dif.append(f'fuente col {c}')
        if (ca.fill.fill_type, ca.fill.start_color.rgb) != (cb.fill.fill_type, cb.fill.start_color.rgb):
            dif.append(f'relleno col {c}')
    print(f'{etiqueta:34} {"IDÉNTICO" if not dif else str(len(dif)) + " DIFERENCIAS"}')
    for d in dif[:8]:
        print('   ', d)


pk = 'Número de WO'
comparar('FLM/PEXT (con repetidos)', pk, DETALLE, MANUAL, filas_db, True)
comparar('FLM/PEXT (reales)', pk, DETALLE, ['SERVICIO', 'CIUDAD', 'LATITUD'], filas_db, True)
comparar('Genérico (pk repetido)', 'COD_MATERIAL',
         [], ['COD_MATERIAL', 'DESCRIPCION_MATERIAL', 'PROYECTO', 'UM', 'TIPO'], filas_db, False)
comparar('Genérico (sin repetidos)', 'Task Id',
         [], ['ESTADO', 'OBSERVACIONES', 'FECHA ATENDIDO'], filas_db, False)
comparar('Genérico (0 filas)', 'Task Id', [], ['ESTADO', 'OBSERVACIONES'], [], False)
