# -*- coding: utf-8 -*-
"""¿Vale la pena write_only en la plantilla manual? Tiempo, memoria y estilos."""
import io
import json
import os
import sqlite3
import time
import tracemalloc

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB = os.path.join(os.environ['TEMP'], 'bench_nucleus.db')
con = sqlite3.connect(DB)
filas_db = [(k, json.loads(dj)) for k, dj in
            con.execute('select key_value,data_json from nucleus_data where proyecto_id=14')]
DETALLE = ['Fecha de creación (WO Creation date)', 'Nombre de Site', 'Autin TT', 'Departamento',
           'Fault Level', 'Estado de la tarea (WO State)', 'Prioridad del Site', 'Provincia', 'Distrito']
MANUAL = ['SERVICIO', 'CIUDAD', 'TECNICO ASIGNADO', 'CONTRATA', 'LATITUD', 'LONGITUD', 'MOTIVO DE AVERÍA',
          'SOLUCIÓN', 'SISTEMAS', 'INICIO DE PARADA', 'FIN DE PARADA', 'BITACORA', 'DETALLE CORRECTIVO',
          'REQUIERE BIÓTICOS', 'MONTO BIÓTICOS (SOLES)', 'COSTO DE MATERIAL (SOLES)']
COLUMNAS = ['Número de WO'] + DETALLE + MANUAL
FILAS = [[k] + [str(d.get(c) or '') for c in DETALLE + MANUAL] for k, d in filas_db]
FILL_DET = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
FILL_GES = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
ESTILOS = {c: FILL_GES for c in MANUAL}
ESTILOS.update({c: FILL_DET for c in DETALLE})


def normal():
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Plantilla'
    hoja.append(list(COLUMNAS))
    anchos = [len(str(c)) for c in COLUMNAS]
    for fila in FILAS:
        hoja.append(fila)
        for i, v in enumerate(fila):
            if len(v) > anchos[i]:
                anchos[i] = len(v)
    for i, a in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(i)].width = min(a + 4, 38)
    hoja.freeze_panes = 'A2'
    hoja.auto_filter.ref = hoja.dimensions
    fuente = Font(bold=True, size=9)
    for i, c in enumerate(COLUMNAS, 1):
        celda = hoja.cell(row=1, column=i)
        celda.font = fuente
        if c in ESTILOS:
            celda.fill = ESTILOS[c]
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def write_only():
    libro = Workbook(write_only=True)
    hoja = libro.create_sheet('Plantilla')
    anchos = [len(str(c)) for c in COLUMNAS]
    fuente = Font(bold=True, size=9)
    cabecera = []
    for c in COLUMNAS:
        celda = WriteOnlyCell(hoja, value=c)
        celda.font = fuente
        if c in ESTILOS:
            celda.fill = ESTILOS[c]
        cabecera.append(celda)
    hoja.append(cabecera)
    for fila in FILAS:
        hoja.append(fila)
        for i, v in enumerate(fila):
            if len(v) > anchos[i]:
                anchos[i] = len(v)
    for i, a in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(i)].width = min(a + 4, 38)
    hoja.freeze_panes = 'A2'
    hoja.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNAS))}{len(FILAS) + 1}'
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


for etiqueta, fn in (('normal (openpyxl)', normal), ('write_only', write_only)):
    fn()
    t = time.perf_counter()
    for _ in range(3):
        datos = fn()
    ms = (time.perf_counter() - t) / 3 * 1000
    tracemalloc.start()
    fn()
    _, pico = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    print(f'{etiqueta:18} {ms:7.0f} ms   pico {pico / 1048576:6.1f} MB   {len(datos) / 1024:.0f} KB')

# ¿Sobreviven congelado, autofiltro, anchos y estilos?
import openpyxl  # noqa: E402
import zipfile  # noqa: E402
datos = write_only()
hoja = openpyxl.load_workbook(io.BytesIO(datos))['Plantilla']
print('freeze:', hoja.freeze_panes, '| autofiltro:', hoja.auto_filter.ref, '| dims:', hoja.dimensions)
print('ancho A:', hoja.column_dimensions['A'].width, '| ancho B:', hoja.column_dimensions['B'].width)
h = hoja.cell(row=1, column=1)
print('cabecera A bold/size/fill:', h.font.bold, h.font.size, h.fill.start_color.rgb)
h2 = hoja.cell(row=1, column=2)
print('cabecera B fill:', h2.fill.start_color.rgb)
print('fila 2 col 1:', repr(hoja.cell(row=2, column=1).value), '| col 2:', repr(hoja.cell(row=2, column=2).value))
with zipfile.ZipFile(io.BytesIO(datos)) as z:
    print('sheet1 sin sharedStrings:', not any(n == 'xl/sharedStrings.xml' for n in z.namelist()))
